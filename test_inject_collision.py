"""
Focused regression test for the Issue-1 root cause.

The original ``inject_images`` copied source media into the output package
under their ORIGINAL filenames (word/media/imageN.png) and appended them after
the existing output entries.  When the output (template) package already owned
a media file with the same canonical name (e.g. word/media/image1.png for a
logo), this produced a duplicate ZIP entry and the wrong image bytes could end
up rendered for a question — the "image of Q1 shows up in Q2" symptom.

This test crafts that exact collision (source image1.png vs template
image1.png, but with distinct relationship ids as in a real merged document)
and asserts the fixed ``inject_images``:
  • preserves the template's own image bytes,
  • writes the source image under a fresh, unique name,
  • rebinds the source reference to the source bytes,
  • leaves the template reference bound to the template bytes,
  • produces a package with no duplicate ZIP entries.
"""

import io
import struct
import zlib
import zipfile
from copy import deepcopy

from docx import Document
from lxml import etree

import mcq_transfer as M

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
W_NS = M.W_NS
SRC_RID = "rIdSourceImageUnique"   # distinct from any template rId


def _png_chunk(typ, data):
    body = typ + data
    return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)


def make_png(r, g, b):
    sig  = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw  = b"\x00" + bytes([r & 255, g & 255, b & 255])
    return sig + _png_chunk(b"IHDR", ihdr) + _png_chunk(b"IDAT", zlib.compress(raw)) + _png_chunk(b"IEND", b"")


def _doc_with_one_image(path, png):
    doc = Document()
    doc.add_paragraph().add_run().add_picture(io.BytesIO(png))
    doc.save(path)


def _rewrite_source_image_rid(src_path):
    """
    Rename the source image relationship id to SRC_RID so that, when merged
    into the template-based output, it does NOT clash with the template's own
    relationship ids — isolating the media *filename* collision under test.
    """
    with zipfile.ZipFile(src_path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    img_target = None
    for el in rels:
        if el.get("Type") == M.IMG_NS:
            el.set("Id", SRC_RID)
            img_target = el.get("Target")
            break
    assert img_target, "source has no image relationship"
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels, xml_declaration=True,
                                                            encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(src_path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in parts.items():
            z.writestr(n, d)
    return img_target


def _template_image_rid(out_path):
    with zipfile.ZipFile(out_path) as z:
        rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))
    for el in rels:
        if el.get("Type") == M.IMG_NS:
            return el.get("Id")
    raise AssertionError("template has no image relationship")


def test_template_media_name_collision_is_safe(tmp_path):
    bytes_src = make_png(10, 20, 30)     # the question's (source) image
    bytes_tpl = make_png(200, 100, 50)   # the template's own image (a "logo")
    assert bytes_src != bytes_tpl

    src = str(tmp_path / "src.docx")
    out = str(tmp_path / "out.docx")
    _doc_with_one_image(src, bytes_src)   # source: word/media/image1.png = bytes_src
    _doc_with_one_image(out, bytes_tpl)   # output: word/media/image1.png = bytes_tpl (collision!)

    _rewrite_source_image_rid(src)        # source image now uses SRC_RID
    rid_tpl = _template_image_rid(out)    # template image keeps its own rId
    assert rid_tpl != SRC_RID

    # Simulate the writer deep-copying a source image paragraph (referencing the
    # SOURCE rId) into the output body, alongside the template's own image.
    with zipfile.ZipFile(out) as z:
        doc_root = etree.fromstring(z.read("word/document.xml"))
        other = {n: z.read(n) for n in z.namelist() if n != "word/document.xml"}
    body = doc_root.find(f"{{{W_NS}}}body")
    tpl_p = next(p for p in body.findall(f"{{{W_NS}}}p")
                 if any(etree.QName(e).localname == "blip" for e in p.iter()))
    src_p = deepcopy(tpl_p)
    for blip in src_p.iter():
        if etree.QName(blip).localname == "blip":
            blip.set(f"{{{R_NS}}}embed", SRC_RID)   # references the SOURCE image
    sectpr = body.find(f"{{{W_NS}}}sectPr")
    (sectpr.addprevious if sectpr is not None else body.append)(src_p)

    new_doc = etree.tostring(doc_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", new_doc)
        for name, data in other.items():
            z.writestr(name, data)

    # ── run the fixed injector ───────────────────────────────────────
    M.inject_images(src, out)

    # ── verify ───────────────────────────────────────────────────────
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert len(names) == len(set(names)), "duplicate ZIP entries!"
        media = {n: z.read(n) for n in names if n.startswith("word/media/")}
        rels  = {el.get("Id"): el.get("Target")
                 for el in etree.fromstring(z.read("word/_rels/document.xml.rels"))}
        doc_root = etree.fromstring(z.read("word/document.xml"))

    # Template image bytes preserved (NOT overwritten by colliding source img).
    assert media["word/media/image1.png"] == bytes_tpl, (
        "template image was overwritten by the colliding source image"
    )
    # Source image stored under a fresh, unique name.
    src_media = [n for n in media if "mcq_img_" in n]
    assert len(src_media) == 1
    assert media[src_media[0]] == bytes_src

    # Resolve blips → rels → media bytes and confirm both bindings are correct.
    bound = []
    for blip in doc_root.iter():
        if etree.QName(blip).localname != "blip":
            continue
        target = rels[blip.get(f"{{{R_NS}}}embed")]
        bound.append(media["word/" + target.lstrip("/")])

    assert bytes_tpl in bound, "template image reference lost / repointed"
    assert bytes_src in bound, "source image reference not bound to source bytes"


if __name__ == "__main__":
    import pytest
    raise SystemExit(pytest.main([__file__, "-v"]))
