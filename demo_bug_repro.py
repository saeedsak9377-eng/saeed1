"""
Before/after demonstration of the image-desynchronisation root cause.

It reproduces the ORIGINAL ``inject_images`` behaviour (copy source media under
their original filenames, appended after the existing package entries) and the
FIXED behaviour, on a document where the template already owns
``word/media/image1.png`` and the source's question image is ALSO
``word/media/image1.png`` (different bytes).

OLD  → duplicate ZIP entry for image1.png; the question's image is overwritten
       by / confused with the template image  → "wrong image in question".
NEW  → source image stored under a unique name, both references correct.
"""

import io
import os
import struct
import tempfile
import zlib
import zipfile
from copy import deepcopy

from docx import Document
from lxml import etree

import mcq_transfer as M

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
W_NS = M.W_NS
IMG_NS = M.IMG_NS
SRC_RID = "rIdSourceImageUnique"


def _chunk(t, d):
    b = t + d
    return struct.pack(">I", len(d)) + b + struct.pack(">I", zlib.crc32(b) & 0xFFFFFFFF)

def png(r, g, b):
    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return sig + _chunk(b"IHDR", ihdr) + _chunk(b"IDAT", zlib.compress(b"\x00" + bytes([r, g, b]))) + _chunk(b"IEND", b"")


def one_image_doc(path, data):
    d = Document()
    d.add_paragraph().add_run().add_picture(io.BytesIO(data))
    d.save(path)


def build_collision(tmp):
    src = os.path.join(tmp, "src.docx")
    out = os.path.join(tmp, "out.docx")
    b_src = png(10, 20, 30)
    b_tpl = png(200, 100, 50)
    one_image_doc(src, b_src)
    one_image_doc(out, b_tpl)

    # give the source image a distinct rId
    with zipfile.ZipFile(src) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    for el in rels:
        if el.get("Type") == IMG_NS:
            el.set("Id", SRC_RID)
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in parts.items():
            z.writestr(n, d)

    # inject a source-referencing image paragraph into the template output
    with zipfile.ZipFile(out) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        other = {n: z.read(n) for n in z.namelist() if n != "word/document.xml"}
    body = root.find(f"{{{W_NS}}}body")
    tpl_p = next(p for p in body.findall(f"{{{W_NS}}}p")
                 if any(etree.QName(e).localname == "blip" for e in p.iter()))
    sp = deepcopy(tpl_p)
    for bl in sp.iter():
        if etree.QName(bl).localname == "blip":
            bl.set(f"{{{R_NS}}}embed", SRC_RID)
    sect = body.find(f"{{{W_NS}}}sectPr")
    (sect.addprevious if sect is not None else body.append)(sp)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True))
        for n, d in other.items():
            z.writestr(n, d)
    return src, out, b_src, b_tpl


def old_inject_images(source_path, output_path):
    """Faithful reproduction of the ORIGINAL (buggy) injector."""
    with zipfile.ZipFile(source_path) as src_zip:
        rels_name = "word/_rels/document.xml.rels"
        src_rels = etree.fromstring(src_zip.read(rels_name))
        src_img_rels = {el.get("Id"): el.get("Target") for el in src_rels if el.get("Type") == IMG_NS}
        src_media = {n: src_zip.read(n) for n in src_zip.namelist() if n.startswith("word/media/")}
    if not src_img_rels:
        return
    with zipfile.ZipFile(output_path) as out_zip:
        out_rels = etree.fromstring(out_zip.read("word/_rels/document.xml.rels"))
    used_ids = {el.get("Id") for el in out_rels}
    counter, rid_remap = 100, {}
    for old_rid in src_img_rels:
        while f"rId{counter}" in used_ids:
            counter += 1
        rid_remap[old_rid] = f"rId{counter}"
        used_ids.add(f"rId{counter}")
        counter += 1
    fd, tmp = tempfile.mkstemp(suffix=".docx"); os.close(fd)
    with zipfile.ZipFile(output_path) as out_zip:
        doc_root = etree.fromstring(out_zip.read("word/document.xml"))
        for blip in doc_root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
            old = blip.get(f"{{{R_NS}}}embed")
            if old in rid_remap:
                blip.set(f"{{{R_NS}}}embed", rid_remap[old])
        new_doc = etree.tostring(doc_root, xml_declaration=True, encoding="UTF-8", standalone=True)
        out_rels = etree.fromstring(out_zip.read("word/_rels/document.xml.rels"))
        for old_rid, target in src_img_rels.items():
            el = etree.SubElement(out_rels, "Relationship")
            el.set("Id", rid_remap[old_rid]); el.set("Type", IMG_NS); el.set("Target", target)
        new_rels = etree.tostring(out_rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as nz:
            for item in out_zip.infolist():
                f = item.filename
                if f == "word/document.xml":
                    nz.writestr(f, new_doc)
                elif f == "word/_rels/document.xml.rels":
                    nz.writestr(f, new_rels)
                else:
                    nz.writestr(item, out_zip.read(f))
            for media_name, media_bytes in src_media.items():   # ORIGINAL names → collision
                nz.writestr(media_name, media_bytes)
    import shutil
    shutil.move(tmp, output_path)


def resolve_bindings(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        media = {n: z.read(n) for n in names if n.startswith("word/media/")}
        rels = {el.get("Id"): el.get("Target") for el in etree.fromstring(z.read("word/_rels/document.xml.rels"))}
        root = etree.fromstring(z.read("word/document.xml"))
    dup = len(names) != len(set(names))
    bound = []
    for bl in root.iter():
        if etree.QName(bl).localname == "blip":
            tgt = rels.get(bl.get(f"{{{R_NS}}}embed"))
            key = "word/" + (tgt or "").lstrip("/")
            bound.append(media.get(key))
    return dup, bound, media


def main():
    import shutil
    print("=" * 70)
    print("IMAGE DESYNCHRONISATION — BEFORE/AFTER DEMONSTRATION")
    print("=" * 70)

    with tempfile.TemporaryDirectory() as tmp:
        # ---- OLD (buggy) ----
        src, out, b_src, b_tpl = build_collision(tmp)
        old_inject_images(src, out)
        dup, bound, media = resolve_bindings(out)
        print("\n[OLD injector — original behaviour]")
        print(f"  duplicate ZIP entries present : {dup}")
        print(f"  media files in package        : {sorted(os.path.basename(k) for k in media)}")
        print(f"  template image bytes intact   : {media.get('word/media/image1.png') == b_tpl}")
        print(f"  source image present verbatim : {b_src in [v for v in media.values()]}")
        print(f"  both images correctly bound   : {b_src in bound and b_tpl in bound}")
        old_ok = (not dup) and (b_src in bound) and (b_tpl in bound)
        print(f"  >>> RESULT: {'OK' if old_ok else 'CORRUPTED (image desynchronised)'}")

        # ---- NEW (fixed) ----
        src, out, b_src, b_tpl = build_collision(tmp)
        M.inject_images(src, out)
        dup, bound, media = resolve_bindings(out)
        print("\n[NEW injector — fixed behaviour]")
        print(f"  duplicate ZIP entries present : {dup}")
        print(f"  media files in package        : {sorted(os.path.basename(k) for k in media)}")
        print(f"  template image bytes intact   : {media.get('word/media/image1.png') == b_tpl}")
        print(f"  source image present verbatim : {b_src in [v for v in media.values()]}")
        print(f"  both images correctly bound   : {b_src in bound and b_tpl in bound}")
        new_ok = (not dup) and (b_src in bound) and (b_tpl in bound) and media.get('word/media/image1.png') == b_tpl
        print(f"  >>> RESULT: {'OK (image stays bound to its question)' if new_ok else 'CORRUPTED'}")

    print("\n" + "=" * 70)
    print(f"OLD injector correct? {old_ok}    NEW injector correct? {new_ok}")
    print("=" * 70)


if __name__ == "__main__":
    main()
