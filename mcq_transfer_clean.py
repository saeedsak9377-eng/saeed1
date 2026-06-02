"""
MCQ Bank Transfer Tool (clean rewrite)
=====================================

Focused on fixing:
1) media/math ordering drift across questions
2) sequence-number leakage (including attached prefixes like "1.النص")
3) large-batch stability by avoiding stale XML anchors
"""

from __future__ import annotations

import argparse
import logging
import os
import re
import shutil
import tempfile
import zipfile
from copy import deepcopy
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import docx
from lxml import etree


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("mcq_transfer_clean")


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
IMG_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

OPTION_KEYS = ["A", "B", "C", "D"]
OPTION_MAP = {"A": "أ", "B": "ب", "C": "ت", "D": "ث"}
ARABIC_LETTERS = ["أ", "ب", "ت", "ث"]

_RE_Q_HEADER = re.compile(r"^QUESTION\s+(\d+)", re.IGNORECASE)
_RE_OPTION = re.compile(r"^([A-D])\.\s*(.*)", re.IGNORECASE | re.DOTALL)
_RE_OPT_PREFIX = re.compile(r"^[A-D]\.\s*", re.IGNORECASE)
_RE_SEQ_PREFIX = re.compile(r"^(?:[\u0660-\u0669\u06F0-\u06F9]+|\d+)[.\)\-]\s*")
_RE_PURE_SEQ = re.compile(r"^[\s]*(?:[\u0660-\u0669\u06F0-\u06F9]+|\d+)[.\)\-\s،؟]*$")
_RE_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_RE_NYMR_CODE = re.compile(r"\[\[QuestionCode\s*:\s*(.+?)(?:\]\]|$)", re.IGNORECASE)
_RE_NYMR_Q_OPEN = re.compile(r"\[\[Question:\s*(.*)", re.IGNORECASE | re.DOTALL)
_RE_NYMR_CH_OPEN = re.compile(r"\[\[Choices:\s*(.*)", re.IGNORECASE | re.DOTALL)
_RE_NYMR_CLOSE = re.compile(r"\]\]")
_RE_ARABIC_PERIOD = re.compile(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF][^.]*\.")

_PEA_SKIP = frozenset(["REFER TO THE FOLLOWING", "أسئلة الاختيار", "أسئلة المقارنة", "فيما يلي سؤال"])
_SIMPLE_SKIP = frozenset(["أسئلة الاختيار من متعدد", "أسئلة المقارنة", "فيما يلي سؤال", "REFER TO THE FOLLOWING"])
_NYMR_NOISE = frozenset(["سري", "Secret", "سري | Secret", "Secret | سري"])


@dataclass
class QuestionRecord:
    seq: int
    ext_id: str
    correct_key: str
    correct_ar: str
    source_fmt: str

    body_paras: List = field(default_factory=list)
    option_paras: Dict[str, List] = field(default_factory=dict)
    body_text: str = ""
    option_texts: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    def get_option_paras(self, key: str) -> List:
        return self.option_paras.get(key, [])


def _clean_ctrl(text: str) -> str:
    return _RE_CTRL.sub("", text or "").strip()


def _p_has_drawing(p) -> bool:
    return bool(p._element.xpath('.//*[local-name()="drawing"]'))


def _p_has_math(p) -> bool:
    return bool(p._element.xpath('.//*[local-name()="oMath"]'))


def _p_has_content(p) -> bool:
    return bool((p.text or "").strip()) or _p_has_drawing(p) or _p_has_math(p)


def _is_pure_seq_number(text: str) -> bool:
    s = (text or "").strip()
    return bool(s) and bool(_RE_PURE_SEQ.fullmatch(s))


def _is_pure_seq(text: str) -> bool:
    return _is_pure_seq_number(text)


def _skip_pea(text: str) -> bool:
    return any(x in text for x in _PEA_SKIP)


def _skip_simple(text: str) -> bool:
    return any(x in text for x in _SIMPLE_SKIP)


def _find_sectpr(body_el):
    return body_el.find(f"{{{W_NS}}}sectPr")


def _insert_before_sectpr(body_el, new_p, sectpr) -> None:
    if sectpr is not None:
        sectpr.addprevious(new_p)
    else:
        body_el.append(new_p)


def detect_source_format(source_doc) -> str:
    full_text = "\n".join((p.text or "") for p in source_doc.paragraphs)
    if "[[QuestionCode" in full_text or "[[Question:" in full_text:
        return "nymr"
    if source_doc.tables and _RE_Q_HEADER.search(full_text):
        return "pea"
    if len(re.findall(r"(?im)^[A-D]\.\s+", full_text)) >= 4:
        return "simple"
    return "pea" if source_doc.tables else "nymr"


def _extract_exam_key(source_doc) -> Dict[int, Dict]:
    key: Dict[int, Dict] = {}
    if not source_doc.tables:
        return key
    for row in source_doc.tables[0].rows[1:]:
        cells = row.cells
        if len(cells) < 5:
            continue
        seq_text = (cells[0].text or "").strip()
        ext_id = (cells[2].text or "").strip()
        ans = (cells[4].text or "").strip().upper()
        if not seq_text.isdigit():
            continue
        seq = int(seq_text)
        if seq in key:
            continue
        key[seq] = {
            "ext_id": ext_id if ext_id else "UNKNOWN",
            "correct_key": ans if ans in OPTION_KEYS else "",
            "correct_ar": OPTION_MAP.get(ans, ""),
        }
    return key


def _parse_pea_raw(source_doc) -> Dict[int, Dict]:
    raw: Dict[int, Dict] = {}
    current_q: Optional[int] = None
    current_part: Optional[str] = None
    seen_q = set()

    def _valid_next(part: str, letter: str) -> bool:
        if part == "body":
            return letter == "A"
        if part in OPTION_KEYS:
            return OPTION_KEYS.index(letter) > OPTION_KEYS.index(part)
        return False

    for p in source_doc.paragraphs:
        text = (p.text or "").strip()
        has_img = _p_has_drawing(p)
        has_math = _p_has_math(p)
        if not (text or has_img or has_math):
            continue

        qh = _RE_Q_HEADER.match(text)
        if qh:
            q_num = int(qh.group(1))
            if q_num in seen_q:
                current_q = None
                current_part = None
                continue
            seen_q.add(q_num)
            current_q = q_num
            current_part = "body"
            raw[current_q] = {"body": [], "A": [], "B": [], "C": [], "D": []}
            continue

        if _skip_pea(text) or current_q is None:
            continue

        om = _RE_OPTION.match(text) if text else None
        if om:
            letter = om.group(1).upper()
            inline = (om.group(2) or "").strip()
            if current_part is not None and _valid_next(current_part, letter):
                current_part = letter
                if inline or has_img or has_math:
                    raw[current_q][current_part].append(p)
                continue

        if current_part is not None:
            raw[current_q][current_part].append(p)

    return raw


def parse_pea(source_doc) -> Tuple[List[QuestionRecord], Dict]:
    exam_key = _extract_exam_key(source_doc)
    raw = _parse_pea_raw(source_doc)
    records: List[QuestionRecord] = []
    for seq, q_num in enumerate(sorted(raw.keys()), start=1):
        q = raw[q_num]
        k = exam_key.get(q_num, {"ext_id": "UNKNOWN", "correct_key": "", "correct_ar": ""})
        records.append(
            QuestionRecord(
                seq=seq,
                ext_id=k["ext_id"],
                correct_key=k["correct_key"],
                correct_ar=k["correct_ar"],
                source_fmt="pea",
                body_paras=q["body"],
                option_paras={x: q[x] for x in OPTION_KEYS},
            )
        )
    return records, exam_key


def parse_nymr(source_doc) -> List[QuestionRecord]:
    lines = []
    for p in source_doc.paragraphs:
        t = _clean_ctrl(p.text)
        if t and t not in _NYMR_NOISE:
            lines.append((t, p))

    records: List[QuestionRecord] = []
    i = 0
    n = len(lines)
    seen_codes = set()

    while i < n:
        text, _ = lines[i]
        cm = _RE_NYMR_CODE.search(text)
        if not cm:
            i += 1
            continue

        code = cm.group(1).strip()
        if code in seen_codes:
            i += 1
            continue
        seen_codes.add(code)
        i += 1

        stem_parts: List[str] = []
        while i < n:
            t, _ = lines[i]
            qm = _RE_NYMR_Q_OPEN.search(t)
            if qm:
                inline = (qm.group(1) or "").strip()
                if _RE_NYMR_CLOSE.search(inline):
                    stem_parts.append(_RE_NYMR_CLOSE.sub("", inline).strip())
                elif inline:
                    stem_parts.append(inline)
                i += 1
                break
            i += 1

        while i < n:
            t, _ = lines[i]
            if _RE_NYMR_CLOSE.search(t):
                rem = _RE_NYMR_CLOSE.sub("", t).strip()
                if rem:
                    stem_parts.append(rem)
                i += 1
                break
            stem_parts.append(t)
            i += 1

        choices: List[Tuple[str, bool]] = []

        def _parse_choice(raw: str):
            is_ok = "*" in raw
            clean = re.sub(r"\s*\*\s*", " ", raw).strip()
            if clean:
                choices.append((clean, is_ok))

        while i < n:
            t, _ = lines[i]
            cm2 = _RE_NYMR_CH_OPEN.search(t)
            if cm2:
                inline = (cm2.group(1) or "").strip()
                if inline and not _RE_NYMR_CLOSE.search(inline):
                    _parse_choice(inline)
                i += 1
                break
            i += 1

        while i < n:
            t, _ = lines[i]
            if _RE_NYMR_CLOSE.search(t):
                rem = _RE_NYMR_CLOSE.sub("", t).strip()
                if rem:
                    _parse_choice(rem)
                i += 1
                break
            _parse_choice(t)
            i += 1

        idx = next((k for k, (_, ok) in enumerate(choices) if ok), 0)
        key = OPTION_KEYS[idx] if idx < len(OPTION_KEYS) else "A"
        ar = ARABIC_LETTERS[idx] if idx < len(ARABIC_LETTERS) else ""
        records.append(
            QuestionRecord(
                seq=len(records) + 1,
                ext_id=code,
                correct_key=key,
                correct_ar=ar,
                source_fmt="nymr",
                body_text=" ".join(stem_parts).strip(),
                option_texts=[x for x, _ in choices],
            )
        )
    return records


def parse_simple(source_doc) -> List[QuestionRecord]:
    records: List[QuestionRecord] = []
    cur_body: List = []
    cur_opts: Dict[str, List] = {k: [] for k in OPTION_KEYS}
    cur_part: Optional[str] = None
    empty_count = 0

    def _flush():
        nonlocal cur_body, cur_opts, cur_part
        if cur_body or any(cur_opts.values()):
            records.append(
                QuestionRecord(
                    seq=len(records) + 1,
                    ext_id=f"Q{len(records)+1}",
                    correct_key="",
                    correct_ar="",
                    source_fmt="simple",
                    body_paras=cur_body,
                    option_paras=cur_opts,
                )
            )
        cur_body = []
        cur_opts = {k: [] for k in OPTION_KEYS}
        cur_part = "body"

    for p in source_doc.paragraphs:
        text = (p.text or "").strip()
        has_img = _p_has_drawing(p)
        has_math = _p_has_math(p)
        empty = not text and not has_img and not has_math

        if empty:
            empty_count += 1
            continue

        prev_empty = empty_count
        empty_count = 0

        if _skip_simple(text):
            empty_count = prev_empty
            continue

        om = _RE_OPTION.match(text) if text else None
        if om:
            letter = om.group(1).upper()
            inline = (om.group(2) or "").strip()
            if cur_part == "body" and letter == "A":
                cur_part = "A"
                if inline or has_img or has_math:
                    cur_opts["A"].append(p)
                continue
            if cur_part in OPTION_KEYS and OPTION_KEYS.index(letter) > OPTION_KEYS.index(cur_part):
                cur_part = letter
                if inline or has_img or has_math:
                    cur_opts[letter].append(p)
                continue

        start_new = (
            cur_part is None
            or (prev_empty >= 2 and cur_part == "body")
            or (prev_empty >= 2 and cur_part == "D")
            or (prev_empty >= 2 and cur_part in OPTION_KEYS and not om)
        )
        if start_new:
            if cur_part is not None or cur_body or any(cur_opts.values()):
                _flush()
            cur_part = "body"

        if cur_part == "body":
            cur_body.append(p)
        elif cur_part in OPTION_KEYS:
            cur_opts[cur_part].append(p)

    if cur_body or any(cur_opts.values()):
        _flush()
    return records


def parse_source(source_doc) -> Tuple[List[QuestionRecord], str]:
    fmt = detect_source_format(source_doc)
    if fmt == "pea":
        recs, _ = parse_pea(source_doc)
    elif fmt == "nymr":
        recs = parse_nymr(source_doc)
    else:
        recs = parse_simple(source_doc)
    return recs, fmt


def validate_records(records: List[QuestionRecord]) -> List[QuestionRecord]:
    seen = {}
    for rec in records:
        if rec.ext_id in seen:
            rec.warnings.append(f"Duplicate ext_id '{rec.ext_id}' (first seq {seen[rec.ext_id]})")
        else:
            seen[rec.ext_id] = rec.seq
        if rec.ext_id in ("", "UNKNOWN"):
            rec.warnings.append("Missing ext_id")
        if not rec.correct_key:
            rec.warnings.append("No correct answer found")
    return records


def _strip_seq_prefix_xml(p_element) -> None:
    for t_el in p_element.findall(f".//{{{W_NS}}}t"):
        if not t_el.text:
            continue
        cleaned = _RE_SEQ_PREFIX.sub("", t_el.text, count=1)
        if cleaned != t_el.text:
            t_el.text = cleaned
            if not cleaned:
                parent_r = t_el.getparent()
                if parent_r is not None and parent_r.tag == f"{{{W_NS}}}r":
                    parent_p = parent_r.getparent()
                    if parent_p is not None:
                        parent_p.remove(parent_r)
        break


def _strip_opt_prefix_xml(p_element) -> None:
    for t_el in p_element.findall(f".//{{{W_NS}}}t"):
        if not t_el.text:
            continue
        cleaned = _RE_OPT_PREFIX.sub("", t_el.text, count=1)
        if cleaned != t_el.text:
            t_el.text = cleaned
            if not cleaned.strip():
                parent_r = t_el.getparent()
                if parent_r is not None and parent_r.tag == f"{{{W_NS}}}r":
                    parent_p = parent_r.getparent()
                    if parent_p is not None:
                        parent_p.remove(parent_r)
        break


def _set_list_paragraph_style(p_element) -> None:
    pPr = p_element.find(f"{{{W_NS}}}pPr")
    if pPr is None:
        pPr = etree.SubElement(p_element, f"{{{W_NS}}}pPr")
    pStyle = pPr.find(f"{{{W_NS}}}pStyle")
    if pStyle is None:
        pStyle = etree.SubElement(pPr, f"{{{W_NS}}}pStyle")
    pStyle.set(f"{{{W_NS}}}val", "ListParagraph")


def _apply_rtl_ppr(p_element) -> None:
    pPr = p_element.find(f"{{{W_NS}}}pPr")
    if pPr is None:
        pPr = etree.SubElement(p_element, f"{{{W_NS}}}pPr")
        p_element.insert(0, pPr)
    jc = pPr.find(f"{{{W_NS}}}jc")
    if jc is None:
        jc = etree.SubElement(pPr, f"{{{W_NS}}}jc")
    jc.set(f"{{{W_NS}}}val", "right")


def _build_lxml_list_text_para(text: str):
    p = etree.Element(f"{{{W_NS}}}p")
    pPr = etree.SubElement(p, f"{{{W_NS}}}pPr")
    pStyle = etree.SubElement(pPr, f"{{{W_NS}}}pStyle")
    pStyle.set(f"{{{W_NS}}}val", "ListParagraph")
    jc = etree.SubElement(pPr, f"{{{W_NS}}}jc")
    jc.set(f"{{{W_NS}}}val", "right")
    if _RE_ARABIC_PERIOD.search(text):
        rPr = etree.SubElement(pPr, f"{{{W_NS}}}rPr")
        etree.SubElement(rPr, f"{{{W_NS}}}rtl")
    r = etree.SubElement(p, f"{{{W_NS}}}r")
    t = etree.SubElement(r, f"{{{W_NS}}}t")
    t.text = text
    if text and (text[0] == " " or text[-1] == " "):
        t.set(XML_SPACE, "preserve")
    return p


def _add_para(output_doc, text: str, style_name: str) -> None:
    p = output_doc.add_paragraph(style=style_name)
    p.clear()
    p.add_run(text)
    pPr = p._element.find(f"{{{W_NS}}}pPr")
    if pPr is None:
        pPr = etree.SubElement(p._element, f"{{{W_NS}}}pPr")
        p._element.insert(0, pPr)
    jc = pPr.find(f"{{{W_NS}}}jc")
    if jc is None:
        jc = etree.SubElement(pPr, f"{{{W_NS}}}jc")
    jc.set(f"{{{W_NS}}}val", "right")
    if _RE_ARABIC_PERIOD.search(text):
        rPr = pPr.find(f"{{{W_NS}}}rPr")
        if rPr is None:
            rPr = etree.SubElement(pPr, f"{{{W_NS}}}rPr")
        if rPr.find(f"{{{W_NS}}}rtl") is None:
            etree.SubElement(rPr, f"{{{W_NS}}}rtl")


def _write_choice_nymr(output_doc, body_el, choice_paras, letter_prefix: str, is_correct: bool, style_name: str) -> None:
    star = " *" if is_correct else ""
    has_img = any(_p_has_drawing(p) for p in choice_paras)
    has_math = any(_p_has_math(p) for p in choice_paras)

    if not has_img and not has_math:
        parts = []
        for p in choice_paras:
            raw = (p.text or "").strip()
            if raw:
                parts.append(_RE_OPT_PREFIX.sub("", raw, count=1).strip())
        plain = " ".join(x for x in parts if x).strip()
        _add_para(output_doc, f"{letter_prefix}{plain}{star}".strip(), style_name)
        return

    for idx, cp in enumerate(choice_paras):
        cp_img = _p_has_drawing(cp)
        cp_math = _p_has_math(cp)

        if cp_img or cp_math:
            new_p = deepcopy(cp._element)
            _set_list_paragraph_style(new_p)
            _apply_rtl_ppr(new_p)
            _strip_opt_prefix_xml(new_p)

            if idx == 0:
                prefix_r = etree.Element(f"{{{W_NS}}}r")
                prefix_t = etree.SubElement(prefix_r, f"{{{W_NS}}}t")
                prefix_t.text = letter_prefix
                prefix_t.set(XML_SPACE, "preserve")
                first_r = new_p.find(f"{{{W_NS}}}r")
                if first_r is not None:
                    first_r.addprevious(prefix_r)
                else:
                    new_p.append(prefix_r)

            if is_correct and idx == len(choice_paras) - 1:
                star_r = etree.Element(f"{{{W_NS}}}r")
                star_t = etree.SubElement(star_r, f"{{{W_NS}}}t")
                star_t.text = " *"
                star_t.set(XML_SPACE, "preserve")
                new_p.append(star_r)

            sectpr = _find_sectpr(body_el)
            _insert_before_sectpr(body_el, new_p, sectpr)
            continue

        plain = _RE_OPT_PREFIX.sub("", (cp.text or "").strip(), count=1).strip()
        line = f"{letter_prefix}{plain}".strip() if idx == 0 else plain
        if is_correct and idx == len(choice_paras) - 1:
            line = f"{line}{star}".strip()
        text_p = _build_lxml_list_text_para(line)
        sectpr = _find_sectpr(body_el)
        _insert_before_sectpr(body_el, text_p, sectpr)


def write_nymr_output(template_doc, output_path: str, records: List[QuestionRecord]) -> None:
    output_doc = docx.Document()
    ns = W_NS
    try:
        tmpl_styles = template_doc.part.styles._element
        out_styles = output_doc.part.styles._element
        for style_el in tmpl_styles.findall(f"{{{ns}}}style"):
            sid = style_el.get(f"{{{ns}}}styleId")
            if not out_styles.findall(f'.//{{{ns}}}style[@{{{ns}}}styleId="{sid}"]'):
                out_styles.append(deepcopy(style_el))
    except Exception as exc:
        log.warning("Style copy warning: %s", exc)

    style_name = "List Paragraph"
    try:
        output_doc.styles[style_name]
    except KeyError:
        output_doc.styles.add_style(style_name, docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)

    body_el = output_doc.element.body
    for p in body_el.findall(f"{{{ns}}}p"):
        body_el.remove(p)

    for rec in records:
        sectpr = _find_sectpr(body_el)
        _add_para(output_doc, f"[[QuestionCode : {rec.ext_id}]]", style_name)
        _add_para(output_doc, " [[Question: ", style_name)

        if rec.source_fmt == "nymr":
            if rec.body_text.strip():
                _add_para(output_doc, rec.body_text.strip(), style_name)
        else:
            for bp in rec.body_paras:
                bp_text = (bp.text or "").strip()
                bp_img = _p_has_drawing(bp)
                bp_math = _p_has_math(bp)

                if not bp_text and not bp_img and not bp_math:
                    continue
                if (not bp_img and not bp_math) and _is_pure_seq_number(bp_text):
                    continue

                if bp_img or bp_math:
                    new_p = deepcopy(bp._element)
                    _strip_seq_prefix_xml(new_p)
                    _set_list_paragraph_style(new_p)
                    _apply_rtl_ppr(new_p)
                    sectpr = _find_sectpr(body_el)
                    _insert_before_sectpr(body_el, new_p, sectpr)
                    continue

                text = _RE_SEQ_PREFIX.sub("", bp_text, count=1).strip()
                if not text:
                    continue
                _add_para(output_doc, text, style_name)

        _add_para(output_doc, "]]", style_name)
        _add_para(output_doc, "[[Choices:", style_name)

        for idx, key in enumerate(OPTION_KEYS):
            prefix = f"{key}. "
            is_correct = (key == rec.correct_key)
            if rec.source_fmt == "nymr":
                opt = rec.option_texts[idx] if idx < len(rec.option_texts) else ""
                _add_para(output_doc, f"{prefix}{opt}{' *' if is_correct else ''}".strip(), style_name)
            else:
                _write_choice_nymr(output_doc, body_el, rec.get_option_paras(key), prefix, is_correct, style_name)

        _add_para(output_doc, "]]", style_name)
        sep = etree.Element(f"{{{ns}}}p")
        sectpr = _find_sectpr(body_el)
        _insert_before_sectpr(body_el, sep, sectpr)

    output_doc.save(output_path)


def inject_images(source_path: str, output_path: str) -> None:
    with zipfile.ZipFile(source_path, "r") as src_zip:
        rels_name = "word/_rels/document.xml.rels"
        if rels_name not in src_zip.namelist():
            return
        src_rels = etree.fromstring(src_zip.read(rels_name))
        src_img_rels = {
            el.get("Id"): el.get("Target")
            for el in src_rels
            if el.get("Type") == IMG_REL_TYPE
        }
        src_media = {
            n: src_zip.read(n)
            for n in src_zip.namelist()
            if n.startswith("word/media/")
        }

    if not src_img_rels:
        return

    with zipfile.ZipFile(output_path, "r") as out_zip:
        out_rels = etree.fromstring(out_zip.read("word/_rels/document.xml.rels"))
    used_ids = {el.get("Id") for el in out_rels}

    rid_remap = {}
    counter = 100
    for old_rid in src_img_rels:
        while f"rId{counter}" in used_ids:
            counter += 1
        new_rid = f"rId{counter}"
        rid_remap[old_rid] = new_rid
        used_ids.add(new_rid)
        counter += 1

    fd, tmp_path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        with zipfile.ZipFile(output_path, "r") as out_zip:
            doc_root = etree.fromstring(out_zip.read("word/document.xml"))
            r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            for blip in doc_root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                old = blip.get(f"{{{r_ns}}}embed")
                if old in rid_remap:
                    blip.set(f"{{{r_ns}}}embed", rid_remap[old])
            new_doc_xml = etree.tostring(doc_root, xml_declaration=True, encoding="UTF-8", standalone=True)

            out_rels = etree.fromstring(out_zip.read("word/_rels/document.xml.rels"))
            for old_rid, target in src_img_rels.items():
                rel = etree.SubElement(out_rels, "Relationship")
                rel.set("Id", rid_remap[old_rid])
                rel.set("Type", IMG_REL_TYPE)
                rel.set("Target", target)
            new_rels_xml = etree.tostring(out_rels, xml_declaration=True, encoding="UTF-8", standalone=True)

            with zipfile.ZipFile(tmp_path, "w", compression=zipfile.ZIP_DEFLATED) as nz:
                for item in out_zip.infolist():
                    name = item.filename
                    if name == "word/document.xml":
                        nz.writestr(name, new_doc_xml)
                    elif name == "word/_rels/document.xml.rels":
                        nz.writestr(name, new_rels_xml)
                    else:
                        nz.writestr(item, out_zip.read(name))
                for media_name, media_bytes in src_media.items():
                    nz.writestr(media_name, media_bytes)
        shutil.move(tmp_path, output_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise


def process_and_transfer(source_path: str, template_path: str, output_path: str) -> None:
    source_doc = docx.Document(source_path)
    template_doc = docx.Document(template_path)
    records, fmt = parse_source(source_doc)
    if not records:
        raise ValueError(f"No questions parsed. Detected format: {fmt}")
    validate_records(records)
    write_nymr_output(template_doc, output_path, records)
    inject_images(source_path, output_path)
    log.info("Transfer complete: %d questions", len(records))


def main() -> int:
    parser = argparse.ArgumentParser(description="MCQ Transfer Clean Script")
    parser.add_argument("--source", required=True, help="Source .docx")
    parser.add_argument("--template", required=True, help="Nymr template .docx")
    parser.add_argument("--output", required=True, help="Output .docx")
    args = parser.parse_args()
    process_and_transfer(args.source, args.template, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
