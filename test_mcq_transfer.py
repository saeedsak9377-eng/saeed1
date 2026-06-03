"""
End-to-end tests for mcq_transfer.py focusing on the two reported bugs:

  Issue 1 — Image / Equation desynchronisation
      Each question's image / equation must remain bound to THAT question's
      block in the output (no leaking into a neighbouring question).

  Issue 2 — Large batches (> 50 questions)
      The pipeline must scale to large mixed batches with zero cross-question
      data bleed (text, images, equations, answer keys).

The tests build synthetic PEA source documents in which every question owns a
uniquely-coloured PNG (so the exact bytes can be traced through the whole
pipeline) plus a uniquely-tokenised OMML equation, run the real
``process_and_transfer`` pipeline against the generated templates, then crack
open the resulting .docx ZIP and assert that every image / equation / text
fragment ended up in — and only in — its own question.
"""

import io
import struct
import zlib
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from lxml import etree

import mcq_transfer as M

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


# ──────────────────────────────────────────────────────────────────────
# Helpers — build unique media + equations
# ──────────────────────────────────────────────────────────────────────
def _png_chunk(typ: bytes, data: bytes) -> bytes:
    body = typ + data
    return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)


def make_png(r: int, g: int, b: int) -> bytes:
    """Return a valid 1x1 RGB PNG with the given colour (unique bytes)."""
    sig  = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)   # 1x1, 8-bit, RGB
    raw  = b"\x00" + bytes([r & 255, g & 255, b & 255])    # filter + pixel
    idat = zlib.compress(raw)
    return sig + _png_chunk(b"IHDR", ihdr) + _png_chunk(b"IDAT", idat) + _png_chunk(b"IEND", b"")


def add_image_paragraph(doc, png_bytes: bytes):
    p = doc.add_paragraph()
    p.add_run().add_picture(io.BytesIO(png_bytes))
    return p


def add_equation_paragraph(doc, token: str):
    """Add a paragraph containing a minimal OMML equation carrying `token`."""
    p = doc.add_paragraph()
    omath = parse_xml(
        f'<m:oMath xmlns:m="{M_NS}">'
        f'<m:r><m:t>{token}</m:t></m:r>'
        f'</m:oMath>'
    )
    p._p.append(omath)
    return p


# ──────────────────────────────────────────────────────────────────────
# Source document builders
# ──────────────────────────────────────────────────────────────────────
def build_pea_source(path, n_questions, *, with_images=True, with_equations=True):
    """Build a PEA-format source with N questions; return per-question meta."""
    doc = Document()

    # EXAMKEY table (col0=seq, col2=ext_id, col4=answer); row0 is header.
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Seq", "x", "ExtID"
    hdr[3].text, hdr[4].text = "y", "Ans"

    meta = {}
    answers = ["A", "B", "C", "D"]
    for q in range(1, n_questions + 1):
        ext_id = f"TSQ{q:05d}"
        ans    = answers[q % 4]
        row = table.add_row().cells
        row[0].text = str(q)
        row[2].text = ext_id
        row[4].text = ans

        png = make_png(q & 255, (q * 7) & 255, (q * 13) & 255) if with_images else None
        tok = f"EQTOKEN{q:05d}" if with_equations else None
        meta[ext_id] = {
            "body_text": f"BodyText Q{q} unique",
            "body_png":  png,
            "eq_token":  tok,
            "answer":    ans,
        }

    for q in range(1, n_questions + 1):
        ext_id = f"TSQ{q:05d}"
        info   = meta[ext_id]
        doc.add_paragraph(f"QUESTION {q} REFER TO THE FOLLOWING")
        doc.add_paragraph(info["body_text"])
        if info["body_png"] is not None:
            add_image_paragraph(doc, info["body_png"])
        if info["eq_token"] is not None:
            add_equation_paragraph(doc, info["eq_token"])
        doc.add_paragraph(f"A. optA Q{q}")
        doc.add_paragraph(f"B. optB Q{q}")
        doc.add_paragraph(f"C. optC Q{q}")
        doc.add_paragraph(f"D. optD Q{q}")

    doc.save(path)
    return meta


# ──────────────────────────────────────────────────────────────────────
# Output inspectors
# ──────────────────────────────────────────────────────────────────────
def read_zip_parts(path):
    with zipfile.ZipFile(path, "r") as z:
        names = z.namelist()
        doc_xml = z.read("word/document.xml")
        rels    = z.read("word/_rels/document.xml.rels")
        media   = {n: z.read(n) for n in names if n.startswith("word/media/")}
    rels_map = {el.get("Id"): el.get("Target")
                for el in etree.fromstring(rels)}
    return etree.fromstring(doc_xml), rels_map, media, names


def blip_bytes_in(el, rels_map, media):
    """Return the list of media-byte blobs referenced by blips under `el`."""
    out = []
    for blip in el.iter():
        if etree.QName(blip).localname != "blip":
            continue
        rid = blip.get(qn("r:embed")) or blip.get(f"{{{R_NS}}}embed")
        if not rid:
            continue
        target = rels_map.get(rid)
        assert target is not None, f"blip rId {rid} has no relationship"
        key = "word/" + target.lstrip("/")
        assert key in media, f"media {key} missing for rId {rid}"
        out.append(media[key])
    return out


def cell_text(tc):
    """All text under an element, including WordML (<w:t>) and Math (<m:t>)."""
    return "".join(
        el.text or ""
        for el in tc.iter()
        if etree.QName(el).localname == "t"
    )


# ──────────────────────────────────────────────────────────────────────
# TABLE-output tests (Arabic QBank table template, key='pea')
# ──────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("n_questions", [3, 60, 120])
def test_table_image_equation_binding(tmp_path, n_questions):
    src = str(tmp_path / "src.docx")
    out = str(tmp_path / "out.docx")
    meta = build_pea_source(src, n_questions)

    M.process_and_transfer(src, out, template_key="pea")

    out_doc = Document(out)
    _, rels_map, media, _ = read_zip_parts(out)

    assert len(out_doc.tables) == n_questions, (
        f"expected {n_questions} tables, got {len(out_doc.tables)}"
    )

    all_referenced = []
    seen_ext_ids = set()

    for table in out_doc.tables:
        struct = M._discover_template_structure(table)
        id_tc   = struct["id_tc"]      # raw lxml <w:tc> elements
        body_tc = struct["body_tc"]
        assert id_tc is not None and body_tc is not None

        ext_id = cell_text(id_tc).strip()
        assert ext_id in meta, f"unknown ext_id in output: {ext_id!r}"
        seen_ext_ids.add(ext_id)
        info = meta[ext_id]

        body_txt = cell_text(body_tc)

        # ── Text must be THIS question's body, never another's ──────
        assert info["body_text"] in body_txt, (
            f"{ext_id}: body text missing/bled. got={body_txt!r}"
        )

        # ── Equation token must be this question's, and only this one ─
        if info["eq_token"]:
            assert info["eq_token"] in body_txt, (
                f"{ext_id}: equation token missing from its own body; "
                f"got={body_txt!r}"
            )
            for other_id, other in meta.items():
                if other_id != ext_id and other["eq_token"]:
                    assert other["eq_token"] not in body_txt, (
                        f"{ext_id}: equation {other['eq_token']} bled in "
                        f"from {other_id}"
                    )

        # ── Image bytes in this body cell must equal this question's ─
        imgs = blip_bytes_in(body_tc, rels_map, media)
        assert len(imgs) == 1, f"{ext_id}: expected 1 image, got {len(imgs)}"
        assert imgs[0] == info["body_png"], (
            f"{ext_id}: WRONG IMAGE bound to question (desynchronised!)"
        )
        all_referenced.append(imgs[0])

    assert seen_ext_ids == set(meta.keys())

    assert len(all_referenced) == n_questions
    assert len({bytes(b) for b in all_referenced}) == n_questions, (
        "two questions ended up pointing at the same media bytes"
    )

    expected_pngs = {bytes(info["body_png"]) for info in meta.values()}
    present_pngs  = {bytes(b) for b in media.values()}
    assert expected_pngs <= present_pngs


def test_table_answer_key_not_bled(tmp_path):
    """Each table's correct-answer star must match its own EXAMKEY answer."""
    src = str(tmp_path / "src.docx")
    out = str(tmp_path / "out.docx")
    meta = build_pea_source(src, 30, with_images=False, with_equations=False)

    M.process_and_transfer(src, out, template_key="pea")
    out_doc = Document(out)

    for table in out_doc.tables:
        struct = M._discover_template_structure(table)
        ext_id = cell_text(struct["id_tc"]).strip()
        want_ar = M.OPTION_MAP[meta[ext_id]["answer"]]
        starred = [opt["label"] for opt in struct["options"]
                   if opt["star_tc"] is not None and "*" in cell_text(opt["star_tc"])]
        assert starred == [want_ar], (
            f"{ext_id}: starred {starred}, expected [{want_ar}]"
        )


# ──────────────────────────────────────────────────────────────────────
# NYMR-output tests (text template, key='nymr')
# ──────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("n_questions", [3, 75])
def test_nymr_image_order_binding(tmp_path, n_questions):
    src = str(tmp_path / "src.docx")
    out = str(tmp_path / "out.docx")
    meta = build_pea_source(src, n_questions, with_equations=False)

    M.process_and_transfer(src, out, template_key="nymr")

    doc_root, rels_map, media, _ = read_zip_parts(out)
    body = doc_root.find(qn("w:body"))

    current_ext = None
    seen = set()
    for p in body.findall(qn("w:p")):
        txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
        if "[[QuestionCode" in txt:
            cm = M._RE_NYMR_CODE.search(txt)
            current_ext = cm.group(1).strip() if cm else None
        blips = blip_bytes_in(p, rels_map, media)
        for b in blips:
            assert current_ext in meta, "image found before any QuestionCode"
            assert bytes(b) == bytes(meta[current_ext]["body_png"]), (
                f"{current_ext}: image desynchronised in nymr output"
            )
            seen.add(current_ext)

    assert seen == set(meta.keys())
    assert len({bytes(v) for v in media.values()}) == n_questions


# ──────────────────────────────────────────────────────────────────────
# Output integrity — produced docx must re-open and have no duplicate
# zip entries (the old code could create duplicate word/media/imageN.png).
# ──────────────────────────────────────────────────────────────────────
def test_output_package_integrity(tmp_path):
    src = str(tmp_path / "src.docx")
    out = str(tmp_path / "out.docx")
    build_pea_source(src, 80)
    M.process_and_transfer(src, out, template_key="pea")

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert len(names) == len(set(names)), "duplicate entries in output zip"

    Document(out)


# ──────────────────────────────────────────────────────────────────────
# Guard for the collision-proof media-rename logic in inject_images:
# source media must be renamed to the mcq_img_* scheme so they can never
# collide with template media that shares a canonical name like image1.png.
# ──────────────────────────────────────────────────────────────────────
def test_inject_images_media_renamed_uniquely(tmp_path):
    src = str(tmp_path / "src.docx")
    out = str(tmp_path / "out.docx")
    build_pea_source(src, 2)
    M.process_and_transfer(src, out, template_key="pea")

    with zipfile.ZipFile(out) as z:
        out_media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert out_media, "no media injected"
    assert all("mcq_img_" in n for n in out_media), (
        f"source media not renamed to collision-proof scheme: {out_media}"
    )


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
