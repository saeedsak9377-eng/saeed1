#!/usr/bin/env python3
"""
inspect_docs.py — Template Structure Inspector
===============================================
Run this script on your Windows machine to print the exact structure of
both your Source File and Question Bank Template.

Usage (Windows Command Prompt or PowerShell):
    python inspect_docs.py "E:\\Source File Template.docx" "E:\\Question Bank Template.docx"

OR edit the two paths below and just run:
    python inspect_docs.py
"""

import sys
import os

# ── Edit these if you prefer not to pass command-line arguments ──────────────
DEFAULT_SOURCE   = r"E:\Source File Template.docx"
DEFAULT_TEMPLATE = r"E:\Question Bank Template.docx"
# ─────────────────────────────────────────────────────────────────────────────

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    print("ERROR: python-docx is not installed.")
    print("Run:  pip install python-docx")
    sys.exit(1)


def _cell_text(tc_elem) -> str:
    """Extract all text from a <w:tc> XML element."""
    return ''.join(
        (t.text or '') for t in tc_elem.findall('.//' + qn('w:t'))
    ).strip()


def inspect_table(table, label: str, max_rows: int = 30):
    """Print every cell's text for up to max_rows rows of a table."""
    print(f"\n{'='*70}")
    print(f"  {label}")
    print(f"{'='*70}")
    rows_xml = table._tbl.findall(qn('w:tr'))
    total    = len(rows_xml)
    print(f"  Total rows: {total}")
    for r_idx, row_xml in enumerate(rows_xml[:max_rows]):
        cells_xml  = row_xml.findall(qn('w:tc'))
        cell_texts = [f"[{_cell_text(c)!r}]" for c in cells_xml]
        print(f"  Row {r_idx:>3} ({len(cells_xml)} cols): {' '.join(cell_texts)}")
    if total > max_rows:
        print(f"  ... ({total - max_rows} more rows not shown)")


def inspect_paragraphs(doc, max_paras: int = 60):
    """Print the first max_paras paragraphs from the document body."""
    print(f"\n{'='*70}")
    print("  BODY PARAGRAPHS  (outside tables)")
    print(f"{'='*70}")
    count = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        print(f"  Para {count:>3}: {text!r}")
        count += 1
        if count >= max_paras:
            print(f"  ... (stopped at {max_paras} paragraphs)")
            break


def inspect_document(path: str, role: str):
    print(f"\n{'#'*70}")
    print(f"#  {role}")
    print(f"#  {path}")
    print(f"{'#'*70}")

    if not os.path.exists(path):
        print(f"\n  *** FILE NOT FOUND: {path} ***\n")
        return

    doc = Document(path)

    print(f"\n  Tables found: {len(doc.tables)}")

    for i, table in enumerate(doc.tables):
        inspect_table(table, f"TABLE {i}", max_rows=40)

    inspect_paragraphs(doc, max_paras=80)


def main():
    args = sys.argv[1:]
    source_path   = args[0] if len(args) > 0 else DEFAULT_SOURCE
    template_path = args[1] if len(args) > 1 else DEFAULT_TEMPLATE

    print("MCQ Transfer Tool – Document Inspector")
    print("=" * 70)

    inspect_document(source_path,   "SOURCE FILE  (PEA Format)")
    inspect_document(template_path, "QUESTION BANK TEMPLATE  (Arabic)")

    print(f"\n{'='*70}")
    print("  Copy and share the output above so the script can be fine-tuned.")
    print(f"{'='*70}\n")


if __name__ == "__main__":
    main()
