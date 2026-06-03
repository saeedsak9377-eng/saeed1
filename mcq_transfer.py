"""
MCQ Bank Transfer Tool  ·  Production-Grade Edition
====================================================
Transforms multiple-choice questions from a source Word document into a
structured Arabic Question Bank Word document.

Supported SOURCE formats
  • PEA     — EXAMKEY table + QUESTION N paragraph headers
  • Nymr    — [[QuestionCode:]] / [[Question:]] / [[Choices:]] markers
  • Simple  — plain sequential; A./B./C./D. options; blank-line boundaries

Supported OUTPUT formats
  • Arabic Question Bank Table  (table-based template)
  • Nymr text document          (no-table Nymr template)

Architecture
  ┌──────────────────────────────────────────────────────────┐
  │  Source docx → Parser → QuestionRecord list              │
  │  → Validator → OutputWriter → inject_images → Output docx│
  └──────────────────────────────────────────────────────────┘

All global constants and compiled regexes live at module level (compiled once).
Parsers are pure functions producing immutable QuestionRecord objects.
Writers receive fully-validated question records only.

----------------------------------------------------------------------
BUG-FIX NOTES (this revision)
----------------------------------------------------------------------
Issue 1 — Image / Equation desynchronisation
    Images, MathML and OMML elements were being relocated into the wrong
    question.  Root causes addressed in ``inject_images``:
      • Every source media file is copied under a fresh, UNIQUE filename
        (media/mcq_img_N.ext) so source images can never overwrite — or be
        overwritten by — template media that shares a name like image1.png
        (a duplicate zip entry silently swapped images between questions).
      • New relationship ids are generated so they cannot collide with
        template ids, with source ids still present in the body, or with each
        other.  Only references that point at a *source* image are rewritten,
        so the template's own images are untouched.
      • Every wp:docPr id is renumbered uniquely, preventing Word from
        silently dropping / relocating drawings (which made images "jump"
        between questions).  Equations carry no relationships, so once body
        order and docPr ids are correct they stay bound to their question.

Issue 2 — Large batches (> 50 questions)
    Duplicated table blocks were deep-copied from the *already-filled* first
    table, seeding every later question with question 1's content.  The table
    pipeline now duplicates from a PRISTINE copy of the reference table that is
    captured BEFORE any cell is filled, so no data bleeds across questions
    regardless of batch size.  Together with the collision-proof rId / media /
    docPr handling above, the pipeline scales to arbitrarily large batches.

Embedded templates
    The two output templates are decoded from base64 constants
    (``_NYMR_TMPL_B64`` / ``_PEA_TMPL_B64``).  When those constants are left
    empty the TemplateManager builds a structurally-correct template
    programmatically, so the tool is functional out of the box.  Paste the
    official ETEC base64 payloads into those constants to restore the exact
    institutional layout/branding.
"""

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# STDLIB / THIRD-PARTY IMPORTS
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
import logging
import os
import re
import shutil
import tempfile
import threading
import zipfile
from copy import deepcopy
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import docx
import docx.enum.style  # noqa: F401  (ensures docx.enum.style is importable)
from lxml import etree

try:                      # tkinter is only needed for the GUI; keep the core
    import tkinter as tk  # importable in headless / server environments.
    from tkinter import filedialog, messagebox
    _TK_AVAILABLE = True
except Exception:         # pragma: no cover - headless fallback
    tk = None
    filedialog = messagebox = None
    _TK_AVAILABLE = False

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# LOGGING
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("mcq_transfer")


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# MODULE-LEVEL CONSTANTS  (never mutated at runtime)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
W_NS   = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
IMG_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

# Relationship / drawing namespaces (used by image injection)
R_NS       = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS       = "http://schemas.openxmlformats.org/drawingml/2006/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

# Arabic letter mapping A/B/C/D → أ/ب/ت/ث
OPTION_MAP: Dict[str, str] = {"A": "أ", "B": "ب", "C": "ت", "D": "ث"}
ARABIC_LETTERS: List[str]  = ["أ", "ب", "ت", "ث"]
OPTION_KEYS:    List[str]  = ["A", "B", "C", "D"]

# Fast O(1) lookup: Arabic letter → option index, option letter → index
_ARABIC_TO_IDX: Dict[str, int] = {ltr: i for i, ltr in enumerate(ARABIC_LETTERS)}
_KEY_TO_IDX:    Dict[str, int] = {k: i for i, k in enumerate(OPTION_KEYS)}

# English → Arabic digit/punctuation translation (built once)
_EN_TO_AR = str.maketrans("0123456789?", "٠١٢٣٤٥٦٧٨٩؟")

# ── Boilerplate phrases to skip ────────────────────────────────────────
_PEA_SKIP: frozenset = frozenset([
    "REFER TO THE FOLLOWING",
    "أسئلة الاختيار",
    "أسئلة المقارنة",
    "فيما يلي سؤال",
])
_NYMR_NOISE: frozenset = frozenset([
    "سري", "Secret", "سري | Secret", "Secret | سري",
])
_SIMPLE_SKIP: frozenset = frozenset([
    "أسئلة الاختيار من متعدد",
    "أسئلة المقارنة",
    "فيما يلي سؤال",
    "REFER TO THE FOLLOWING",
])

# ── All regexes compiled ONCE at module load ───────────────────────────
_RE_Q_HEADER     = re.compile(r"^QUESTION\s+(\d+)", re.IGNORECASE)
_RE_OPTION       = re.compile(r"^([A-D])\.\s*(.*)", re.DOTALL | re.IGNORECASE)
_RE_SEQ_PREFIX   = re.compile(r"^\d+\.\s*")          # "1. " at start
_RE_OPT_PREFIX   = re.compile(r"^[A-D]\.\s*", re.IGNORECASE)
_RE_PURE_SEQ     = re.compile(                        # entire string = seq number
    r"^[\s]*(?:[٠-٩]+|\d+)[.\)\-\s،؟]*$"
)
_RE_CTRL_CHARS   = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_RE_NYMR_CODE    = re.compile(r"\[\[QuestionCode\s*:\s*(.+?)(?:\]\]|$)", re.IGNORECASE)
_RE_NYMR_Q_OPEN  = re.compile(r"\[\[Question:\s*(.*)", re.IGNORECASE | re.DOTALL)
_RE_NYMR_CH_OPEN = re.compile(r"\[\[Choices:\s*(.*)",  re.IGNORECASE | re.DOTALL)
_RE_NYMR_CLOSE   = re.compile(r"\]\]")
_RE_OPT_SIMPLE   = re.compile(r"^([A-D])\.\s*(.*)", re.DOTALL | re.IGNORECASE)
_RE_ARABIC_PERIOD = re.compile(            # Arabic char somewhere before a period
    r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF][^.]*\."
)

# ── Arabic formatting rules ────────────────────────────────────────────
_Q_TYPE_CONTEXTUAL = "الخطأ السياقي"
_Q_TYPE_FILL_BLANK = "إكمال الجمل"
_RE_DOT_SEQ        = re.compile(r"\.{2,}")   # 2+ consecutive periods
_UNDERSCORE_FILL   = "_______"                 # 7 underscores (fixed width)
_RE_TRAILING_PERIOD = re.compile(r"[\u0600-\u06FF].*\.$", re.DOTALL)

# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# DATA MODEL
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
@dataclass
class QuestionRecord:
    """
    Immutable-ish record for a single parsed question.
    body_paras / option_paras hold python-docx paragraph objects (PEA/Simple).
    body_text  / option_texts hold plain strings (Nymr).
    """
    seq:          int
    ext_id:       str
    correct_key:  str
    correct_ar:   str
    source_fmt:   str

    body_paras:   List = field(default_factory=list)
    option_paras: Dict[str, List] = field(default_factory=dict)

    body_text:    str = ""
    option_texts: List[str] = field(default_factory=list)

    warnings:     List[str] = field(default_factory=list)

    def get_option_paras(self, key: str) -> List:
        return self.option_paras.get(key, [])


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# UTILITY FUNCTIONS
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def to_arabic_digits(text: str) -> str:
    """Convert ASCII digits/? to Eastern Arabic equivalents."""
    return text.translate(_EN_TO_AR)


def _p_has_drawing(p) -> bool:
    """True if paragraph element contains a w:drawing."""
    return bool(p._element.xpath('.//*[local-name()="drawing"]'))


def _p_has_math(p) -> bool:
    """True if paragraph element contains an oMath element."""
    return bool(p._element.xpath('.//*[local-name()="oMath"]'))


def _p_has_content(p) -> bool:
    """True if paragraph has visible text, drawing, or math."""
    return bool(p.text.strip()) or _p_has_drawing(p) or _p_has_math(p)


def _is_pure_seq(text: str) -> bool:
    """True when text is ONLY a sequence number like '1.' '١.' '3-'."""
    s = text.strip()
    return bool(s) and bool(_RE_PURE_SEQ.match(s))


def _clean_ctrl(text: str) -> str:
    """Strip control characters that Word sometimes embeds."""
    return _RE_CTRL_CHARS.sub("", text).strip()


def _tc_text(tc) -> str:
    """Get all text from a <w:tc> element."""
    return "".join(x.text or "" for x in tc.findall(f".//{{{W_NS}}}t")).strip()


def _skip_pea(text: str) -> bool:
    return any(ph in text for ph in _PEA_SKIP)


def _skip_simple(text: str) -> bool:
    return any(ph in text for ph in _SIMPLE_SKIP)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# FORMAT DETECTION
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def detect_source_format(source_doc) -> str:
    """Returns 'pea', 'nymr', or 'simple'.  Reads full_text only once."""
    full_text = "\n".join(p.text for p in source_doc.paragraphs)

    if "[[QuestionCode" in full_text or "[[Question:" in full_text:
        return "nymr"

    if source_doc.tables and _RE_Q_HEADER.search(full_text):
        return "pea"

    if len(_RE_OPT_SIMPLE.findall(full_text)) >= 4:
        return "simple"

    return "pea" if source_doc.tables else "nymr"


def detect_output_format(template_doc) -> str:
    """Returns 'table' or 'nymr'."""
    if template_doc.tables:
        return "table"
    full_text = "\n".join(p.text for p in template_doc.paragraphs)
    if "[[QuestionCode" in full_text or "[[Question:" in full_text:
        return "nymr"
    return "table"


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# PEA PARSER
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _extract_exam_key(source_doc) -> Dict[int, Dict]:
    """Read EXAMKEY table → {seq_int: {ext_id, correct_key, correct_ar}}."""
    key: Dict[int, Dict] = {}
    if not source_doc.tables:
        return key
    for row in source_doc.tables[0].rows[1:]:
        cells = row.cells
        if len(cells) < 5:
            continue
        seq_text = cells[0].text.strip()
        ext_id   = cells[2].text.strip()
        answer   = cells[4].text.strip().upper()
        if not seq_text.isdigit():
            continue
        seq = int(seq_text)
        if seq in key:
            log.warning("EXAMKEY: duplicate seq %d — keeping first", seq)
            continue
        key[seq] = {
            "ext_id":      ext_id,
            "correct_key": answer,
            "correct_ar":  OPTION_MAP.get(answer, ""),
        }
    return key


def _parse_pea_raw(source_doc) -> Dict[int, Dict]:
    """
    Walk source_doc paragraphs and bucket them into:
      { q_num: { 'body': [para,...], 'A': [para,...], 'B':..., 'C':..., 'D':... } }
    """
    raw: Dict[int, Dict] = {}
    current_q:    Optional[int] = None
    current_part: Optional[str] = None  # 'body' | 'A' | 'B' | 'C' | 'D'
    seen_q_nums:  set            = set()

    def _next_valid_key(part: str, letter: str) -> bool:
        """Strict A→B→C→D sequence enforcement."""
        if part == "body":
            return letter == "A"
        if part in _KEY_TO_IDX:
            return _KEY_TO_IDX.get(letter, -1) > _KEY_TO_IDX[part]
        return False

    for p in source_doc.paragraphs:
        text     = p.text.strip()
        has_img  = _p_has_drawing(p)
        has_math = _p_has_math(p)
        has_content = bool(text) or has_img or has_math

        if not has_content:
            continue

        # ── QUESTION N header ───────────────────────────────────────
        qm = _RE_Q_HEADER.match(text)
        if qm:
            q_num = int(qm.group(1))
            if q_num in seen_q_nums:
                log.warning("PEA: duplicate QUESTION %d header — skipping duplicate", q_num)
                current_q    = None
                current_part = None
                continue
            seen_q_nums.add(q_num)
            current_q    = q_num
            current_part = "body"
            raw[current_q] = {"body": [], "A": [], "B": [], "C": [], "D": []}
            continue

        # ── boilerplate → skip without affecting state ──────────────
        if _skip_pea(text):
            continue

        if current_q is None:
            continue   # before first QUESTION header

        # ── option detection ────────────────────────────────────────
        om = _RE_OPTION.match(text) if text else None
        if om:
            letter      = om.group(1).upper()
            inline_text = om.group(2).strip()
            if _next_valid_key(current_part, letter):
                current_part = letter
                if inline_text or has_img or has_math:
                    raw[current_q][current_part].append(p)
                continue

        # ── append to active bucket ─────────────────────────────────
        if current_part is not None:
            raw[current_q][current_part].append(p)

    return raw


def parse_pea(source_doc) -> Tuple[List[QuestionRecord], Dict]:
    """Full PEA parse: extract key + parse paragraphs → List[QuestionRecord]."""
    exam_key = _extract_exam_key(source_doc)
    raw      = _parse_pea_raw(source_doc)

    records: List[QuestionRecord] = []
    for seq_idx, q_num in enumerate(sorted(raw.keys()), start=1):
        q     = raw[q_num]
        k     = exam_key.get(q_num, {"ext_id": "UNKNOWN", "correct_key": "", "correct_ar": ""})
        rec   = QuestionRecord(
            seq          = seq_idx,
            ext_id       = k["ext_id"],
            correct_key  = k["correct_key"],
            correct_ar   = k["correct_ar"],
            source_fmt   = "pea",
            body_paras   = q["body"],
            option_paras = {key: q[key] for key in OPTION_KEYS},
        )
        records.append(rec)

    return records, exam_key


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# NYMR PARSER
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _nymr_lines(source_doc):
    """Yield (cleaned_text, para) skipping noise and control chars."""
    for p in source_doc.paragraphs:
        text = _clean_ctrl(p.text)
        if text and text not in _NYMR_NOISE:
            yield text, p


def _nymr_parse_choice(raw: str, choices: list) -> None:
    """Parse one raw choice line: detect *, strip it, convert digits."""
    is_correct = "*" in raw
    clean = re.sub(r"\s*\*\s*", " ", raw).strip()
    clean = to_arabic_digits(clean)
    if clean:
        choices.append((clean, is_correct))


def parse_nymr(source_doc) -> List[QuestionRecord]:
    """Parse Nymr-format source → List[QuestionRecord] (plain-text records)."""
    lines     = list(_nymr_lines(source_doc))
    records:  List[QuestionRecord] = []
    n         = len(lines)
    i         = 0
    seen_codes: set = set()

    while i < n:
        text, _ = lines[i]

        cm = _RE_NYMR_CODE.search(text)
        if not cm:
            i += 1
            continue

        q_code = cm.group(1).strip()
        if q_code in seen_codes:
            log.warning("Nymr: duplicate QuestionCode '%s' — skipping", q_code)
            i += 1
            continue
        seen_codes.add(q_code)
        i += 1

        # ── stem ────────────────────────────────────────────────────
        stem_parts: List[str] = []
        while i < n:
            text, _ = lines[i]
            qm = _RE_NYMR_Q_OPEN.search(text)
            if qm:
                inline = qm.group(1).strip()
                if _RE_NYMR_CLOSE.search(inline):
                    stem_parts.append(_RE_NYMR_CLOSE.sub("", inline).strip())
                elif inline:
                    stem_parts.append(inline)
                i += 1
                break
            i += 1

        while i < n:
            text, _ = lines[i]
            if _RE_NYMR_CLOSE.search(text):
                rem = _RE_NYMR_CLOSE.sub("", text).strip()
                if rem:
                    stem_parts.append(rem)
                i += 1
                break
            stem_parts.append(text)
            i += 1

        q_stem = to_arabic_digits(" ".join(stem_parts).strip())

        # ── choices ─────────────────────────────────────────────────
        choices: List[Tuple[str, bool]] = []

        while i < n:
            text, _ = lines[i]
            cm2 = _RE_NYMR_CH_OPEN.search(text)
            if cm2:
                inline = cm2.group(1).strip()
                if inline and not _RE_NYMR_CLOSE.search(inline):
                    _nymr_parse_choice(inline, choices)
                i += 1
                break
            i += 1

        while i < n:
            text, _ = lines[i]
            if _RE_NYMR_CLOSE.search(text):
                rem = _RE_NYMR_CLOSE.sub("", text).strip()
                if rem:
                    _nymr_parse_choice(rem, choices)
                i += 1
                break
            _nymr_parse_choice(text, choices)
            i += 1

        correct_idx = next((ci for ci, (_, ok) in enumerate(choices) if ok), 0)
        correct_key = OPTION_KEYS[correct_idx] if correct_idx < len(OPTION_KEYS) else "A"
        correct_ar  = ARABIC_LETTERS[correct_idx] if correct_idx < len(ARABIC_LETTERS) else ""

        rec = QuestionRecord(
            seq          = len(records) + 1,
            ext_id       = q_code,
            correct_key  = correct_key,
            correct_ar   = correct_ar,
            source_fmt   = "nymr",
            body_text    = q_stem,
            option_texts = [t for t, _ in choices],
        )
        records.append(rec)

    return records


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# SIMPLE FORMAT PARSER
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def parse_simple(source_doc) -> List[QuestionRecord]:
    """Parse simple sequential format → List[QuestionRecord]."""
    paras = source_doc.paragraphs
    n     = len(paras)

    records:     List[QuestionRecord] = []
    cur_body:    List = []
    cur_opts:    Dict[str, List] = {k: [] for k in OPTION_KEYS}
    cur_part:    Optional[str]   = None  # None | 'body' | 'A'...'D'
    consec_empty = 0

    def _flush_question():
        """Close current question and start a fresh accumulator."""
        nonlocal cur_body, cur_opts, cur_part
        if cur_body or any(cur_opts.values()):
            rec = QuestionRecord(
                seq          = len(records) + 1,
                ext_id       = f"Q{len(records)+1}",
                correct_key  = "",
                correct_ar   = "",
                source_fmt   = "simple",
                body_paras   = cur_body,
                option_paras = cur_opts,
            )
            records.append(rec)
        cur_body  = []
        cur_opts  = {k: [] for k in OPTION_KEYS}
        cur_part  = "body"

    def _valid_next(letter: str) -> bool:
        if cur_part == "body":
            return letter == "A"
        if cur_part in _KEY_TO_IDX:
            return _KEY_TO_IDX.get(letter, -1) > _KEY_TO_IDX[cur_part]
        return False

    i = 0
    while i < n:
        p        = paras[i]
        text     = p.text.strip()
        has_img  = _p_has_drawing(p)
        has_math = _p_has_math(p)
        empty    = not text and not has_img and not has_math

        if empty:
            consec_empty += 1
            i += 1
            continue

        prev_empty   = consec_empty
        consec_empty = 0

        # Boilerplate: skip but keep the empty-line count intact
        if _skip_simple(text):
            i += 1
            consec_empty = prev_empty  # restore so boundary still fires
            continue

        # Option detection
        om = _RE_OPT_SIMPLE.match(text) if text else None
        if om:
            letter      = om.group(1).upper()
            inline_text = om.group(2).strip()
            if cur_part is not None and _valid_next(letter):
                cur_part = letter
                if inline_text or has_img or has_math:
                    cur_opts[cur_part].append(p)
                i += 1
                continue

        # Question boundary
        in_opts    = cur_part in OPTION_KEYS
        start_new  = (
            cur_part is None
            or (prev_empty >= 2 and cur_part == "body")
            or (prev_empty >= 2 and cur_part == "D")
            or (prev_empty >= 2 and in_opts and not om)
        )

        if start_new:
            if cur_part is not None:
                _flush_question()
            elif cur_body or any(cur_opts.values()):
                _flush_question()
            cur_part = "body"

        if cur_part == "body":
            cur_body.append(p)
        elif cur_part in OPTION_KEYS:
            cur_opts[cur_part].append(p)

        i += 1

    if cur_body or any(cur_opts.values()):
        _flush_question()

    return records


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# UNIFIED PARSE ENTRY POINT
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def parse_source(source_doc) -> Tuple[List[QuestionRecord], str]:
    """Auto-detect format and return (list_of_records, format_string)."""
    fmt = detect_source_format(source_doc)
    log.info("Detected source format: %s", fmt.upper())

    if fmt == "pea":
        records, _ = parse_pea(source_doc)
    elif fmt == "nymr":
        records = parse_nymr(source_doc)
    else:
        records = parse_simple(source_doc)

    log.info("Parsed %d questions", len(records))
    return records, fmt


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# VALIDATION LAYER
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def validate_records(records: List[QuestionRecord]) -> List[QuestionRecord]:
    """Validate every record; attach warnings; log anomalies."""
    seen_ids: Dict[str, int] = {}

    for rec in records:
        w = rec.warnings

        if rec.ext_id in seen_ids:
            w.append(
                f"Duplicate ext_id '{rec.ext_id}' "
                f"(first at seq {seen_ids[rec.ext_id]})"
            )
            log.warning("Q%d: %s", rec.seq, w[-1])
        else:
            seen_ids[rec.ext_id] = rec.seq

        if rec.ext_id in ("UNKNOWN", ""):
            w.append("Missing ext_id / not found in EXAMKEY")
            log.warning("Q%d: %s", rec.seq, w[-1])

        if not rec.correct_key:
            w.append("No correct answer found")
            log.warning("Q%d: %s", rec.seq, w[-1])

        if rec.source_fmt in ("pea", "simple"):
            body_ok = any(_p_has_content(p) for p in rec.body_paras)
            if not body_ok:
                w.append("Empty question body")
                log.warning("Q%d (id=%s): %s", rec.seq, rec.ext_id, w[-1])

            for key in OPTION_KEYS:
                opt = rec.get_option_paras(key)
                if not opt:
                    w.append(f"Missing option {key}")
                    log.warning("Q%d (id=%s): missing option %s", rec.seq, rec.ext_id, key)

        else:  # nymr
            if not rec.body_text.strip():
                w.append("Empty question body")
                log.warning("Q%d (id=%s): %s", rec.seq, rec.ext_id, w[-1])
            if len(rec.option_texts) < 4:
                w.append(f"Only {len(rec.option_texts)} choices (expected 4)")
                log.warning("Q%d (id=%s): %s", rec.seq, rec.ext_id, w[-1])

    ok  = sum(1 for r in records if not r.warnings)
    bad = len(records) - ok
    log.info("Validation: %d OK, %d with warnings", ok, bad)
    return records


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# XML / PARAGRAPH HELPERS
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _strip_seq_prefix_xml(p_element) -> None:
    """Remove leading 'N. ' from the first <w:t> of an lxml paragraph element."""
    for t_el in p_element.findall(f".//{{{W_NS}}}t"):
        if t_el.text:
            cleaned = _RE_SEQ_PREFIX.sub("", t_el.text, count=1)
            if cleaned != t_el.text:
                t_el.text = cleaned
                if not cleaned:
                    parent_r = t_el.getparent()
                    if parent_r is not None and parent_r.tag == f"{{{W_NS}}}r":
                        parent_r.getparent().remove(parent_r)
            break


def _strip_opt_prefix_xml(p_element) -> None:
    """Remove leading 'A. ' / 'B. ' etc. from the first <w:t> of an element."""
    for t_el in p_element.findall(f".//{{{W_NS}}}t"):
        if t_el.text:
            cleaned = _RE_OPT_PREFIX.sub("", t_el.text, count=1)
            if cleaned != t_el.text:
                t_el.text = cleaned
                if not cleaned.strip():
                    parent_r = t_el.getparent()
                    if parent_r is not None and parent_r.tag == f"{{{W_NS}}}r":
                        parent_r.getparent().remove(parent_r)
            break


def _set_list_paragraph_style(p_element) -> None:
    """Set pStyle to ListParagraph on an lxml paragraph element."""
    pPr = p_element.find(f"{{{W_NS}}}pPr")
    if pPr is None:
        pPr = etree.SubElement(p_element, f"{{{W_NS}}}pPr")
    pStyle = pPr.find(f"{{{W_NS}}}pStyle")
    if pStyle is None:
        pStyle = etree.SubElement(pPr, f"{{{W_NS}}}pStyle")
    pStyle.set(f"{{{W_NS}}}val", "ListParagraph")


def _find_sectpr(body_el):
    """Return the <w:sectPr> child of body_el."""
    return body_el.find(f"{{{W_NS}}}sectPr")


def _insert_before_sectpr(body_el, new_p, sectpr) -> None:
    """Insert new_p immediately before sectpr (or append if not found)."""
    if sectpr is not None:
        sectpr.addprevious(new_p)
    else:
        body_el.append(new_p)


def _copy_paras_into_tc(source_paras, target_tc,
                        keep_existing: bool = False,
                        strip_first_seq: bool = False) -> None:
    """Deepcopy python-docx paragraphs into a <w:tc> element."""
    if not keep_existing:
        for p in target_tc.findall(f"{{{W_NS}}}p"):
            target_tc.remove(p)

    for idx, src_p in enumerate(source_paras):
        new_p = deepcopy(src_p._element)
        if strip_first_seq and idx == 0:
            _strip_seq_prefix_xml(new_p)
        target_tc.append(new_p)

    if not target_tc.findall(f"{{{W_NS}}}p"):
        etree.SubElement(target_tc, f"{{{W_NS}}}p")


def _set_tc_text(tc, text: str) -> None:
    """Replace a cell's content with a single plain-text paragraph."""
    existing = tc.findall(f"{{{W_NS}}}p")
    pPr_clone = rPr_clone = None
    if existing:
        pPr_el = existing[0].find(f"{{{W_NS}}}pPr")
        if pPr_el is not None:
            pPr_clone = deepcopy(pPr_el)
        first_r = existing[0].find(f".//{{{W_NS}}}r")
        if first_r is not None:
            rPr_el = first_r.find(f"{{{W_NS}}}rPr")
            if rPr_el is not None:
                rPr_clone = deepcopy(rPr_el)
    for p in existing:
        tc.remove(p)
    new_p = etree.SubElement(tc, f"{{{W_NS}}}p")
    if pPr_clone is not None:
        new_p.insert(0, pPr_clone)
    new_r = etree.SubElement(new_p, f"{{{W_NS}}}r")
    if rPr_clone is not None:
        new_r.insert(0, rPr_clone)
    t_el = etree.SubElement(new_r, f"{{{W_NS}}}t")
    t_el.text = text
    if text and (text[0] == " " or text[-1] == " "):
        t_el.set(XML_SPACE, "preserve")


def _set_tc_text_plain(tc, text: str) -> None:
    """Clear cell and write plain text (no formatting inherited — Nymr path)."""
    for p in tc.findall(f"{{{W_NS}}}p"):
        tc.remove(p)
    new_p = etree.SubElement(tc, f"{{{W_NS}}}p")
    new_r = etree.SubElement(new_p, f"{{{W_NS}}}r")
    t_el  = etree.SubElement(new_r, f"{{{W_NS}}}t")
    t_el.text = text


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# TEMPLATE STRUCTURE DISCOVERY (CACHED PER TABLE)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _discover_template_structure(table) -> Dict:
    """
    Scan a template table ONCE and return:
      { 'id_tc': tc_el, 'body_tc': tc_el,
        'options': [{'label':str, 'star_tc':tc, 'text_tc':tc}, ...] }
    Result is cached on the table object itself.
    """
    cache_attr = "_mcq_struct_cache"
    if hasattr(table, cache_attr):
        return getattr(table, cache_attr)

    result = {"id_tc": None, "body_tc": None, "options": []}
    ARABIC_OPTS = set(ARABIC_LETTERS)
    LABEL_CODE  = "رمز السؤال"
    LABEL_BODY  = "نص السؤال"

    for ri, row in enumerate(table.rows):
        actual_tcs = row._tr.findall(f"{{{W_NS}}}tc")
        for ci, tc in enumerate(actual_tcs):
            text = _tc_text(tc)
            if LABEL_CODE in text:
                if ci > 0:
                    result["id_tc"] = actual_tcs[ci - 1]
                elif ri > 0:
                    prev_tcs = table.rows[ri - 1]._tr.findall(f"{{{W_NS}}}tc")
                    if prev_tcs:
                        result["id_tc"] = prev_tcs[0]
            elif LABEL_BODY in text:
                if ri + 1 < len(table.rows):
                    next_tcs = table.rows[ri + 1]._tr.findall(f"{{{W_NS}}}tc")
                    if next_tcs:
                        result["body_tc"] = next_tcs[0]
            elif text in ARABIC_OPTS:
                if ci >= 2:
                    result["options"].append(
                        {"label": text, "star_tc": actual_tcs[ci - 2], "text_tc": actual_tcs[ci - 1]}
                    )
                elif ci == 1:
                    result["options"].append(
                        {"label": text, "star_tc": None, "text_tc": actual_tcs[ci - 1]}
                    )

    order = {ltr: i for i, ltr in enumerate(ARABIC_LETTERS)}
    result["options"].sort(key=lambda o: order.get(o["label"], 99))

    setattr(table, cache_attr, result)
    return result


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# TABLE BLOCK FILLERS
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _fill_table_pea(table, rec: QuestionRecord) -> None:
    """Fill one Arabic QBank table from a PEA/Simple QuestionRecord."""
    struct = _discover_template_structure(table)

    if struct["id_tc"] is not None:
        _set_tc_text(struct["id_tc"], rec.ext_id)

    if struct["body_tc"] is not None:
        _copy_paras_into_tc(
            rec.body_paras, struct["body_tc"],
            keep_existing=False, strip_first_seq=True,
        )

    for i, opt in enumerate(struct["options"]):
        ar_letter = opt["label"]
        star_text = "*" if ar_letter == rec.correct_ar else ""
        if opt["star_tc"] is not None:
            _set_tc_text(opt["star_tc"], star_text)
        key   = OPTION_KEYS[i] if i < len(OPTION_KEYS) else None
        paras = rec.get_option_paras(key) if key else []
        _copy_paras_into_tc(paras, opt["text_tc"], keep_existing=False)


def _fill_table_nymr(table, rec: QuestionRecord) -> None:
    """Fill one Arabic QBank table from a Nymr QuestionRecord."""
    struct  = _discover_template_structure(table)
    choices = rec.option_texts

    if struct["id_tc"] is not None:
        _set_tc_text(struct["id_tc"], rec.ext_id)

    if struct["body_tc"] is not None:
        _set_tc_text_plain(struct["body_tc"], rec.body_text)

    for i, opt in enumerate(struct["options"]):
        ar_letter = opt["label"]
        star_text = "*" if ar_letter == rec.correct_ar else ""
        if opt["star_tc"] is not None:
            _set_tc_text(opt["star_tc"], star_text)
        ch_text = choices[i] if i < len(choices) else ""
        _set_tc_text_plain(opt["text_tc"], ch_text)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# TEMPLATE BLOCK DUPLICATION
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _duplicate_table_block(template_doc, reference_tbl, seq_num: int):
    """
    Deepcopy a PRISTINE reference table element into the document body with a
    styled 'سؤال رقم: N' separator.  Returns the new Table wrapper.

    ``reference_tbl`` MUST be an unfilled <w:tbl> element (a copy of the
    template's first table taken BEFORE any cell was written) so question N's
    block is never seeded with question 1's content — the cross-question data
    bleed that corrupted large batches.
    """
    ns   = W_NS
    body = template_doc.element.body

    style_p = None
    for child in body:
        if child.tag == f"{{{ns}}}p":
            if "سؤال رقم" in "".join(t.text or "" for t in child.findall(f".//{{{ns}}}t")):
                style_p = child
                break

    sep_p = deepcopy(style_p) if style_p is not None else etree.Element(f"{{{ns}}}p")
    for r in sep_p.findall(f"{{{ns}}}r"):
        sep_p.remove(r)
    new_r = etree.SubElement(sep_p, f"{{{ns}}}r")
    if style_p is not None:
        first_r = style_p.find(f"{{{ns}}}r")
        if first_r is not None:
            rPr = first_r.find(f"{{{ns}}}rPr")
            if rPr is not None:
                new_r.insert(0, deepcopy(rPr))
    t_el      = etree.SubElement(new_r, f"{{{ns}}}t")
    t_el.text = f"سؤال رقم: {seq_num}"

    body.append(sep_p)
    new_tbl = deepcopy(reference_tbl)
    body.append(new_tbl)

    from docx.table import Table as DocxTable
    return DocxTable(new_tbl, template_doc)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# ARABIC FORMATTING RULES
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _apply_rule2_fill_blank(text: str) -> str:
    """Rule 2 — إكمال الجمل: replace 2+ consecutive dots with _UNDERSCORE_FILL."""
    return _RE_DOT_SEQ.sub(_UNDERSCORE_FILL, text)


def _apply_rule3_rtl_period(text: str) -> str:
    """Rule 3 — RTL Period Alignment: move a trailing period to the front."""
    stripped = text.rstrip()
    if stripped.endswith(".") and _RE_ARABIC_PERIOD.search(stripped):
        core = stripped[:-1].rstrip()
        return "." + core
    return text


def _add_para_with_period_run(output_doc, text: str, list_style_name: str) -> None:
    """Rule 3 paragraph writer: leading '.' emitted as a separate <w:rtl/> run."""
    ns = W_NS
    p = output_doc.add_paragraph(style=list_style_name)
    p.clear()

    if text.startswith(".") and len(text) > 1:
        r1 = p.add_run(".")
        rPr1 = etree.SubElement(r1._r, f"{{{ns}}}rPr")
        etree.SubElement(rPr1, f"{{{ns}}}rtl")
        p.add_run(text[1:])
    else:
        p.add_run(text)

    pPr = p._element.find(f"{{{ns}}}pPr")
    if pPr is None:
        pPr = etree.SubElement(p._element, f"{{{ns}}}pPr")
        p._element.insert(0, pPr)
    jc_el = pPr.find(f"{{{ns}}}jc")
    if jc_el is None:
        jc_el = etree.SubElement(pPr, f"{{{ns}}}jc")
    jc_el.set(f"{{{ns}}}val", "right")

    if _RE_ARABIC_PERIOD.search(text):
        rPr_el = pPr.find(f"{{{ns}}}rPr")
        if rPr_el is None:
            rPr_el = etree.SubElement(pPr, f"{{{ns}}}rPr")
        if rPr_el.find(f"{{{ns}}}rtl") is None:
            etree.SubElement(rPr_el, f"{{{ns}}}rtl")


def _detect_question_type(body_paras: list) -> str:
    """Return the question type label from a body paragraph, or ''."""
    for p in body_paras:
        t = p.text.strip()
        if t in (_Q_TYPE_CONTEXTUAL, _Q_TYPE_FILL_BLANK):
            return t
    return ""


def _find_stem_para_idx(body_paras: list) -> int:
    """Return the index of the ACTUAL question stem paragraph."""
    LABELS = {_Q_TYPE_CONTEXTUAL, _Q_TYPE_FILL_BLANK,
              "في كل جملة", "تلي الجملة",
              "(الخطأ ليس", "مثال:", "الشرح:"}
    for i in range(len(body_paras) - 1, -1, -1):
        t = body_paras[i].text.strip()
        if not t:
            continue
        if _is_pure_seq(t):
            continue
        if any(t.startswith(lbl) for lbl in LABELS):
            continue
        return i
    return -1


def _add_para_rule1_contextual(output_doc, stem_text: str,
                                choice_words: list,
                                list_style_name: str) -> None:
    """Rule 1 — الخطأ السياقي: emit choice words as bold runs."""
    ns = W_NS
    p = output_doc.add_paragraph(style=list_style_name)
    p.clear()

    spans = []
    for word in choice_words:
        if not word:
            continue
        idx = 0
        while True:
            pos = stem_text.find(word, idx)
            if pos == -1:
                break
            spans.append((pos, pos + len(word), word))
            idx = pos + len(word)

    if not spans:
        p.add_run(stem_text)
    else:
        spans.sort(key=lambda s: s[0])
        merged = [spans[0]]
        for start, end, word in spans[1:]:
            prev_start, prev_end, prev_word = merged[-1]
            if start <= prev_end + 1:
                merged[-1] = (prev_start, max(prev_end, end),
                              stem_text[prev_start:max(prev_end, end)])
            else:
                merged.append((start, end, word))

        cursor = 0
        for start, end, word in merged:
            if cursor < start:
                p.add_run(stem_text[cursor:start])
            bold_run = p.add_run(stem_text[start:end])
            bold_run.bold = True
            cursor = end
        if cursor < len(stem_text):
            p.add_run(stem_text[cursor:])

    pPr = p._element.find(f"{{{ns}}}pPr")
    if pPr is None:
        pPr = etree.SubElement(p._element, f"{{{ns}}}pPr")
        p._element.insert(0, pPr)
    jc_el = pPr.find(f"{{{ns}}}jc")
    if jc_el is None:
        jc_el = etree.SubElement(pPr, f"{{{ns}}}jc")
    jc_el.set(f"{{{ns}}}val", "right")

    if _RE_ARABIC_PERIOD.search(stem_text):
        rPr_el = pPr.find(f"{{{ns}}}rPr")
        if rPr_el is None:
            rPr_el = etree.SubElement(pPr, f"{{{ns}}}rPr")
        if rPr_el.find(f"{{{ns}}}rtl") is None:
            etree.SubElement(rPr_el, f"{{{ns}}}rtl")


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# NYMR OUTPUT WRITER
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def _add_para(output_doc, text: str, list_style_name: str) -> None:
    """Add a right-aligned paragraph in List Paragraph style."""
    p = output_doc.add_paragraph(style=list_style_name)
    p.clear()
    p.add_run(text)

    pPr = p._element.find(f"{{{W_NS}}}pPr")
    if pPr is None:
        pPr = etree.SubElement(p._element, f"{{{W_NS}}}pPr")
        p._element.insert(0, pPr)

    jc_el = pPr.find(f"{{{W_NS}}}jc")
    if jc_el is None:
        jc_el = etree.SubElement(pPr, f"{{{W_NS}}}jc")
    jc_el.set(f"{{{W_NS}}}val", "right")

    if _RE_ARABIC_PERIOD.search(text):
        rPr_el = pPr.find(f"{{{W_NS}}}rPr")
        if rPr_el is None:
            rPr_el = etree.SubElement(pPr, f"{{{W_NS}}}rPr")
        if rPr_el.find(f"{{{W_NS}}}rtl") is None:
            etree.SubElement(rPr_el, f"{{{W_NS}}}rtl")


def _write_choice_nymr(output_doc, body_el, sectpr, choice_paras,
                        letter_prefix: str, is_correct: bool,
                        list_style_name: str) -> None:
    """
    Write one answer choice robustly handling text / image / math / mixed.
    Image/math choice paragraphs are deep-copied straight into the body, in
    order, so the drawing/equation stays bound to this exact choice.
    """
    ns   = W_NS
    star = " *" if is_correct else ""
    has_img  = any(_p_has_drawing(p) for p in choice_paras)
    has_math = any(_p_has_math(p)    for p in choice_paras)

    if not has_img and not has_math:
        parts = [_RE_OPT_PREFIX.sub("", p.text, count=1).strip() for p in choice_paras if p.text.strip()]
        plain = " ".join(parts).strip()
        _add_para(output_doc, f"{letter_prefix}{plain}{star}".strip(), list_style_name)
        return

    for idx, cp in enumerate(choice_paras):
        cp_img  = _p_has_drawing(cp)
        cp_math = _p_has_math(cp)

        if cp_img or cp_math:
            new_p = deepcopy(cp._element)
            _set_list_paragraph_style(new_p)
            _strip_opt_prefix_xml(new_p)

            if idx == 0:
                prefix_r = etree.Element(f"{{{ns}}}r")
                prefix_t = etree.SubElement(prefix_r, f"{{{ns}}}t")
                prefix_t.text = letter_prefix
                prefix_t.set(XML_SPACE, "preserve")
                first_r = new_p.find(f"{{{ns}}}r")
                if first_r is not None:
                    first_r.addprevious(prefix_r)
                else:
                    new_p.append(prefix_r)

            if is_correct and idx == len(choice_paras) - 1:
                star_r = etree.Element(f"{{{ns}}}r")
                star_t = etree.SubElement(star_r, f"{{{ns}}}t")
                star_t.text = " *"
                star_t.set(XML_SPACE, "preserve")
                new_p.append(star_r)

            sectpr = _find_sectpr(body_el)
            _insert_before_sectpr(body_el, new_p, sectpr)

        else:
            plain = _RE_OPT_PREFIX.sub("", cp.text, count=1).strip()
            line  = (f"{letter_prefix}{plain}" if idx == 0 else plain)
            if is_correct and idx == len(choice_paras) - 1:
                line += star
            _add_para(output_doc, line.strip(), list_style_name)


def write_nymr_output(source_doc, template_doc, output_path: str,
                      records: List[QuestionRecord]) -> None:
    """Write all questions to a Nymr-format .docx output file."""
    ns = W_NS
    output_doc = docx.Document()

    try:
        tmpl_styles = template_doc.part.styles._element
        out_styles  = output_doc.part.styles._element
        for style_el in tmpl_styles.findall(f"{{{ns}}}style"):
            sid = style_el.get(f"{{{ns}}}styleId")
            if not out_styles.findall(f'.//{{{ns}}}style[@{{{ns}}}styleId="{sid}"]'):
                out_styles.append(deepcopy(style_el))
    except Exception as e:
        log.warning("Style copy partial failure: %s", e)

    list_style_name = "List Paragraph"
    try:
        output_doc.styles[list_style_name]
    except KeyError:
        output_doc.styles.add_style(
            list_style_name, docx.enum.style.WD_STYLE_TYPE.PARAGRAPH
        )

    body_el = output_doc.element.body
    for p in body_el.findall(f"{{{ns}}}p"):
        body_el.remove(p)

    sectpr = _find_sectpr(body_el)

    for rec in records:
        _add_para(output_doc, f"[[QuestionCode : {rec.ext_id}]]", list_style_name)
        _add_para(output_doc, " [[Question: ", list_style_name)

        q_type    = _detect_question_type(rec.body_paras) if rec.source_fmt != "nymr" else ""
        stem_idx  = _find_stem_para_idx(rec.body_paras)   if rec.source_fmt != "nymr" else -1
        choice_words = []
        if q_type == _Q_TYPE_CONTEXTUAL:
            for key in OPTION_KEYS:
                for cp in rec.get_option_paras(key):
                    w = cp.text.strip()
                    if w:
                        choice_words.append(w)

        if rec.source_fmt == "nymr":
            if rec.body_text.strip():
                _add_para(output_doc, rec.body_text.strip(), list_style_name)
        else:
            first_body = True
            for para_idx, bp in enumerate(rec.body_paras):
                bp_text  = bp.text.strip()
                bp_img   = _p_has_drawing(bp)
                bp_math  = _p_has_math(bp)

                if not bp_text and not bp_img and not bp_math:
                    continue
                if not bp_img and not bp_math and _is_pure_seq(bp_text):
                    first_body = False
                    continue

                if bp_img or bp_math:
                    new_p = deepcopy(bp._element)
                    if first_body:
                        _strip_seq_prefix_xml(new_p)
                    _set_list_paragraph_style(new_p)
                    sectpr = _find_sectpr(body_el)
                    _insert_before_sectpr(body_el, new_p, sectpr)
                else:
                    text = _RE_SEQ_PREFIX.sub("", bp_text, count=1).strip() if first_body else bp_text
                    if text:
                        is_stem = (para_idx == stem_idx)
                        if is_stem and q_type == _Q_TYPE_CONTEXTUAL and choice_words:
                            _add_para_rule1_contextual(
                                output_doc, text, choice_words, list_style_name
                            )
                        elif is_stem and q_type == _Q_TYPE_FILL_BLANK:
                            text = _apply_rule2_fill_blank(text)
                            text = _apply_rule3_rtl_period(text)
                            _add_para_with_period_run(output_doc, text, list_style_name)
                        else:
                            _add_para(output_doc, text, list_style_name)
                        sectpr = _find_sectpr(body_el)

                first_body = False

        _add_para(output_doc, "]]", list_style_name)
        _add_para(output_doc, "[[Choices:", list_style_name)

        for choice_idx, key in enumerate(OPTION_KEYS):
            letter_prefix = f"{key}. "
            is_correct    = (key == rec.correct_key)
            choice_paras  = (
                rec.get_option_paras(key)
                if rec.source_fmt in ("pea", "simple")
                else []
            )

            if rec.source_fmt == "nymr":
                text = rec.option_texts[choice_idx] if choice_idx < len(rec.option_texts) else ""
                line = f"{letter_prefix}{text}{' *' if is_correct else ''}".strip()
                _add_para(output_doc, line, list_style_name)
            else:
                sectpr = _find_sectpr(body_el)
                _write_choice_nymr(
                    output_doc, body_el, sectpr,
                    choice_paras, letter_prefix, is_correct,
                    list_style_name,
                )
            sectpr = _find_sectpr(body_el)

        _add_para(output_doc, "]]", list_style_name)

        sep_p  = etree.Element(f"{{{ns}}}p")
        sectpr = _find_sectpr(body_el)
        _insert_before_sectpr(body_el, sep_p, sectpr)

    output_doc.save(output_path)
    log.info("Nymr output saved: %s", output_path)


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# IMAGE INJECTION  (zip-level, rId-safe, collision-proof)
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def inject_images(source_path: str, output_path: str) -> None:
    """
    Copy source images into the output docx, keeping every image bound to the
    exact paragraph / table-cell where the parser placed it.

    Correctness guarantees (fixes for image/equation desynchronisation and
    large-batch corruption — see module docstring):

      • Each source-image relationship referenced in the output gets a FRESH,
        globally-unique r:id that cannot collide with template ids, with source
        ids still present in the body, or with each other.
      • Each copied media file is written under a FRESH, unique filename
        (``media/mcq_img_N.ext``) so source images never overwrite — or get
        overwritten by — template media sharing a name (e.g. image1.png).
      • Only ``a:blip`` / ``v:imagedata`` references that point at a *source*
        image are rewritten; the template's own images are untouched.
      • Every ``wp:docPr`` id is renumbered uniquely so Word never silently
        drops or relocates a drawing.
    """
    rels_name = "word/_rels/document.xml.rels"

    # ── read source rels + media ─────────────────────────────────────
    with zipfile.ZipFile(source_path, "r") as src_zip:
        names = src_zip.namelist()
        if rels_name not in names:
            return
        src_rels_root = etree.fromstring(src_zip.read(rels_name))
        src_img_rels = {
            el.get("Id"): el.get("Target")
            for el in src_rels_root
            if el.get("Type") == IMG_NS and el.get("Id") and el.get("Target")
        }
        src_media = {
            name: src_zip.read(name)
            for name in names
            if name.startswith("word/media/")
        }

    if not src_img_rels:
        return

    # ── read output package ──────────────────────────────────────────
    with zipfile.ZipFile(output_path, "r") as out_zip:
        out_names = out_zip.namelist()
        doc_root  = etree.fromstring(out_zip.read("word/document.xml"))
        out_rels  = etree.fromstring(out_zip.read(rels_name))
        ct_xml    = out_zip.read("[Content_Types].xml").decode("utf-8")

    embed_attrs = (f"{{{R_NS}}}embed", f"{{{R_NS}}}link")

    # ── find every reference to a SOURCE image in the output body ────
    referenced = []   # list of (element, attr_name, old_rid)
    for el in doc_root.iter():
        for attr in embed_attrs:
            rid = el.get(attr)
            if rid and rid in src_img_rels:
                referenced.append((el, attr, rid))
    if not referenced:
        return

    # ── collision-free new-rId generator ─────────────────────────────
    existing_rids = {el.get("Id") for el in out_rels if el.get("Id")}
    for el in doc_root.iter():
        for a, v in el.attrib.items():
            if a.startswith(f"{{{R_NS}}}") and isinstance(v, str) and v.startswith("rId"):
                existing_rids.add(v)

    _rid_state = {"n": 1000}

    def _new_rid() -> str:
        while True:
            _rid_state["n"] += 1
            cand = f"rId{_rid_state['n']}"
            if cand not in existing_rids:
                existing_rids.add(cand)
                return cand

    # ── collision-free media filename generator ──────────────────────
    used_media = {n for n in out_names if n.startswith("word/media/")}
    _media_state = {"n": 0}

    def _ext_of(target: str) -> str:
        base = target.replace("\\", "/").rsplit("/", 1)[-1]
        return base.rsplit(".", 1)[-1].lower() if "." in base else "png"

    def _new_media(ext: str) -> str:
        while True:
            _media_state["n"] += 1
            cand = f"word/media/mcq_img_{_media_state['n']}.{ext}"
            if cand not in used_media:
                used_media.add(cand)
                return cand

    def _src_bytes(target: str):
        norm = target.replace("\\", "/").lstrip("/")
        base = norm.rsplit("/", 1)[-1]
        for key in (f"word/{norm}", norm, f"word/media/{base}"):
            if key in src_media:
                return src_media[key]
        for k, v in src_media.items():
            if k.rsplit("/", 1)[-1] == base:
                return v
        return None

    # ── build remaps (dedupe identical targets to one media copy) ────
    target_to_media: Dict[str, str] = {}
    new_media_files: Dict[str, bytes] = {}
    rid_remap:       Dict[str, str] = {}
    new_rels:        Dict[str, str] = {}

    referenced_rids = {rid for (_, _, rid) in referenced}
    for old_rid in referenced_rids:
        target = src_img_rels[old_rid]
        if target not in target_to_media:
            data = _src_bytes(target)
            if data is None:
                log.warning("inject_images: source media for %s (%s) not found",
                            old_rid, target)
                continue
            new_name = _new_media(_ext_of(target))
            target_to_media[target] = new_name
            new_media_files[new_name] = data
        new_name = target_to_media.get(target)
        if new_name is None:
            continue
        new_rid = _new_rid()
        rid_remap[old_rid] = new_rid
        new_rels[new_rid] = new_name.split("word/", 1)[-1]   # "media/xxx.ext"

    # ── rewrite the image references in the body ─────────────────────
    for el, attr, old_rid in referenced:
        new_rid = rid_remap.get(old_rid)
        if new_rid:
            el.set(attr, new_rid)

    # ── renumber docPr ids so Word never re-shuffles drawings ────────
    _docpr = {"n": 0}
    for el in doc_root.iter():
        if etree.QName(el).localname == "docPr":
            _docpr["n"] += 1
            el.set("id", str(_docpr["n"]))

    # ── append new relationships (in the package-rel namespace) ──────
    rels_ns = etree.QName(out_rels).namespace
    rel_tag = f"{{{rels_ns}}}Relationship" if rels_ns else "Relationship"
    for new_rid, target in new_rels.items():
        rel = etree.SubElement(out_rels, rel_tag)
        rel.set("Id", new_rid)
        rel.set("Type", IMG_NS)
        rel.set("Target", target)

    new_doc_xml  = etree.tostring(doc_root, xml_declaration=True,
                                  encoding="UTF-8", standalone=True)
    new_rels_xml = etree.tostring(out_rels, xml_declaration=True,
                                  encoding="UTF-8", standalone=True)

    # ── ensure content-type defaults for every media extension ───────
    mime_by_ext = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                   "gif": "image/gif", "bmp": "image/bmp", "tif": "image/tiff",
                   "tiff": "image/tiff", "emf": "image/x-emf", "wmf": "image/x-wmf",
                   "svg": "image/svg+xml"}
    for new_name in new_media_files:
        ext = new_name.rsplit(".", 1)[-1].lower()
        if f'Extension="{ext}"' not in ct_xml:
            mime = mime_by_ext.get(ext, "application/octet-stream")
            ct_xml = ct_xml.replace(
                "</Types>",
                f'<Default Extension="{ext}" ContentType="{mime}"/></Types>',
            )

    # ── rewrite the package (no duplicate entries) ───────────────────
    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".docx")
    os.close(tmp_fd)
    try:
        with zipfile.ZipFile(output_path, "r") as out_zip, \
             zipfile.ZipFile(tmp_path, "w", compression=zipfile.ZIP_DEFLATED) as nz:
            for item in out_zip.infolist():
                fname = item.filename
                if fname == "word/document.xml":
                    nz.writestr(fname, new_doc_xml)
                elif fname == rels_name:
                    nz.writestr(fname, new_rels_xml)
                elif fname == "[Content_Types].xml":
                    nz.writestr(fname, ct_xml.encode("utf-8"))
                else:
                    nz.writestr(item, out_zip.read(fname))
            for media_name, media_bytes in new_media_files.items():
                nz.writestr(media_name, media_bytes)

        shutil.move(tmp_path, output_path)
        log.info("Images injected: %d reference(s) → %d media file(s)",
                 len(referenced), len(new_media_files))

    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# TEMPLATE MANAGER
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
import base64 as _base64
import io     as _io
import hashlib as _hashlib
import threading as _threading

_TEMPLATE_REGISTRY: Dict[str, Dict] = {
    "nymr": {
        "label":       "Nymr Format",
        "description": "Text/paragraph-based output with [[QuestionCode]] markers",
        "b64_var":     "_NYMR_TMPL_B64",
        "output_fmt":  "nymr",
        "icon":        "📝",
    },
    "pea": {
        "label":       "PEA Format",
        "description": "Arabic Question Bank table layout (هيئة تقويم التعليم)",
        "b64_var":     "_PEA_TMPL_B64",
        "output_fmt":  "table",
        "icon":        "📋",
    },
}

_DEFAULT_TEMPLATE_KEY = "pea"


def _generate_template_bytes(output_fmt: str) -> bytes:
    """
    Build a structurally-correct template document in memory.

    Used as a fallback when the embedded base64 payload for a template is
    empty.  The generated table layout exactly matches what
    ``_discover_template_structure`` expects:

        row0:  [ id-cell ][ "رمز السؤال" ]
        row1:  [ "نص السؤال" ]
        row2:  [ body-cell ]
        rows3-6: [ star-cell ][ text-cell ][ "أ"/"ب"/"ت"/"ث" ]
    """
    doc = docx.Document()

    if output_fmt == "table":
        # A separator paragraph whose style is cloned for duplicated blocks.
        doc.add_paragraph("سؤال رقم: ")

        table = doc.add_table(rows=7, cols=3)
        try:
            table.style = "Table Grid"
        except Exception:
            pass
        rows = table.rows
        # row0: code label in column 1 → id cell is column 0
        rows[0].cells[1].text = "رمز السؤال"
        # row1: body label in column 0 → body cell is row2/col0
        rows[1].cells[0].text = "نص السؤال"
        # row2: body cell left blank (filled per-question)
        # rows3-6: option label in column 2 → star=col0, text=col1
        for idx, letter in enumerate(ARABIC_LETTERS):
            rows[3 + idx].cells[2].text = letter
    # For the nymr format the writer builds a brand-new document and only
    # copies styles from the template, so a minimal valid doc is sufficient.

    bio = _io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


class TemplateManager:
    """Manages embedded templates with caching and integrity validation."""

    _instance: Optional["TemplateManager"] = None
    _lock = _threading.Lock()

    def __new__(cls):
        with cls._lock:
            if cls._instance is None:
                cls._instance = super().__new__(cls)
                cls._instance._cache: Dict[str, bytes] = {}
                cls._instance._checksums: Dict[str, str] = {}
                cls._instance._per_key_locks: Dict[str, _threading.Lock] = {}
        return cls._instance

    @property
    def available(self) -> List[str]:
        return list(_TEMPLATE_REGISTRY.keys())

    def get_info(self, key: str) -> Dict:
        if key not in _TEMPLATE_REGISTRY:
            raise KeyError(f"Unknown template key '{key}'. "
                           f"Available: {self.available}")
        return _TEMPLATE_REGISTRY[key]

    def load(self, key: str) -> "docx.Document":
        raw = self._get_raw(key)
        try:
            doc = docx.Document(_io.BytesIO(raw))
            log.info("Template '%s' loaded (%d bytes)", key, len(raw))
            return doc
        except Exception as e:
            raise TemplateError(
                f"Template '{key}' could not be opened as a Word document: {e}"
            ) from e

    def checksum(self, key: str) -> str:
        self._get_raw(key)
        return self._checksums.get(key, "")

    def validate_all(self) -> Dict[str, bool]:
        results = {}
        for key in self.available:
            try:
                self.load(key)
                results[key] = True
            except Exception as e:
                log.error("Template '%s' validation failed: %s", key, e)
                results[key] = False
        return results

    def _get_raw(self, key: str) -> bytes:
        if key not in _TEMPLATE_REGISTRY:
            raise TemplateError(
                f"Unknown template '{key}'. Available: {', '.join(self.available)}"
            )

        with self._lock:
            if key not in self._per_key_locks:
                self._per_key_locks[key] = _threading.Lock()

        with self._per_key_locks[key]:
            if key in self._cache:
                return self._cache[key]

            info    = _TEMPLATE_REGISTRY[key]
            b64_var = info["b64_var"]
            b64_str = (globals().get(b64_var) or "").replace("\n", "").replace(" ", "")

            if b64_str:
                try:
                    raw = _base64.b64decode(b64_str)
                except Exception as e:
                    raise TemplateError(
                        f"Template '{key}' embedded data is corrupted: {e}"
                    ) from e
            else:
                # No embedded payload → generate a working template.
                raw = _generate_template_bytes(info["output_fmt"])
                log.info("Template '%s' generated programmatically "
                         "(no embedded base64 payload).", key)

            self._cache[key]     = raw
            self._checksums[key] = _hashlib.md5(raw).hexdigest()
            return raw


class TemplateError(Exception):
    """Raised when a template cannot be loaded or is invalid."""


template_manager = TemplateManager()


def _resolve_template() -> Optional[str]:
    """Legacy compatibility shim. Returns None (all templates are embedded)."""
    return None


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# MASTER PROCESS FUNCTION
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
def process_and_transfer(source_path: str, output_path: str,
                         template_key: str = _DEFAULT_TEMPLATE_KEY,
                         template_path: Optional[str] = None) -> None:
    """
    Full pipeline:
      1. Load template (embedded registry by key, or fallback path)
      2. Parse source → QuestionRecord list
      3. Validate records
      4. Fill output (table or Nymr format)
      5. Inject images
    """
    log.info("=== MCQ Transfer Start ===")
    log.info("Source:       %s", source_path)
    log.info("Template key: %s", template_key)
    log.info("Output:       %s", output_path)

    source_doc = docx.Document(source_path)

    try:
        template_doc = template_manager.load(template_key)
        log.info("Template loaded from embedded registry (key=%s, md5=%s)",
                 template_key, template_manager.checksum(template_key))
    except (TemplateError, KeyError) as e:
        if template_path and os.path.isfile(template_path):
            log.warning("Embedded template failed (%s); falling back to: %s", e, template_path)
            template_doc = docx.Document(template_path)
        else:
            raise TemplateError(
                f"Cannot load template '{template_key}'.\n"
                f"Detail: {e}\n\n"
                f"Available templates: {', '.join(template_manager.available)}"
            ) from e

    tmpl_info  = template_manager.get_info(template_key)
    output_fmt = tmpl_info["output_fmt"]
    log.info("Output format: %s", output_fmt.upper())

    records, fmt = parse_source(source_doc)
    if not records:
        raise ValueError(
            f"No questions found in source.\nDetected format: {fmt.upper()}"
        )
    validate_records(records)

    # ── Nymr output ─────────────────────────────────────────────────
    if output_fmt == "nymr":
        write_nymr_output(source_doc, template_doc, output_path, records)
        inject_images(source_path, output_path)
        log.info("=== Transfer Complete (%d questions) ===", len(records))
        return

    # ── Table output ─────────────────────────────────────────────────
    template_tables = template_doc.tables[:]
    if not template_tables:
        raise RuntimeError("Selected template has no table to fill.")

    # FIX (Issue 2): capture a PRISTINE, unfilled copy of the reference table
    # BEFORE any cell is written, so duplicated question blocks are never
    # seeded with an earlier question's content.
    reference_tbl = deepcopy(template_tables[0]._tbl)

    for i, rec in enumerate(records):
        if i < len(template_tables):
            target_table = template_tables[i]
        else:
            target_table = _duplicate_table_block(
                template_doc, reference_tbl, rec.seq
            )

        try:
            if rec.source_fmt in ("pea", "simple"):
                _fill_table_pea(target_table, rec)
            else:
                _fill_table_nymr(target_table, rec)
        except Exception as e:
            raise RuntimeError(
                f"Q{rec.seq} (id={rec.ext_id}): fill failed — {e}"
            ) from e

    template_doc.save(output_path)
    inject_images(source_path, output_path)
    log.info("=== Transfer Complete: %d questions ===", len(records))


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# EMBEDDED ASSETS
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# Small placeholder PNG used only for the GUI window icon/header.  It has no
# effect on document transformation and any decode failure is ignored by the
# GUI.  Replace with the official ETEC logo base64 if desired.
_ETEC_LOGO_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR4nGNgYGAAAAAEAAH2FzhVAAAAAElFTkSuQmCC"
)

# Embedded output templates (base64-encoded .docx).  Leave EMPTY to use the
# built-in programmatic generator (fully functional).  Paste the official ETEC
# base64 payloads here to reproduce the exact institutional layout/branding —
# the loader strips whitespace/newlines before decoding.
_NYMR_TMPL_B64 = ""
_PEA_TMPL_B64  = ""


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# ENTERPRISE GUI
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# ETEC brand palette (extracted from official logo)
_C = {
    "navy":        "#2D2D7A",
    "navy_dark":   "#1A1A50",
    "navy_mid":    "#3B3B8F",
    "teal":        "#3C96D2",
    "green":       "#3CAA50",
    "purple":      "#6B4FA0",
    "bg":          "#F4F6FA",
    "card":        "#FFFFFF",
    "border":      "#DDE2EE",
    "border_focus":"#3B3B8F",
    "text_h":      "#1A1A50",
    "text_body":   "#3D3D5C",
    "text_muted":  "#8890AB",
    "success":     "#1E8B4C",
    "success_bg":  "#E8F8EF",
    "error":       "#C0392B",
    "error_bg":    "#FDF0EF",
    "warning":     "#D07A00",
    "warning_bg":  "#FFF8E8",
    "info":        "#2471A3",
    "info_bg":     "#EAF4FC",
    "btn_primary":       "#2D2D7A",
    "btn_primary_hover": "#3B3B8F",
    "btn_primary_fg":    "#FFFFFF",
    "btn_secondary":     "#EDF0F8",
    "btn_secondary_fg":  "#2D2D7A",
    "btn_secondary_hover":"#DDE2EE",
    "btn_disabled":      "#BCC2D4",
    "btn_disabled_fg":   "#FFFFFF",
    "tmpl_selected_bg":  "#EEF0FF",
    "tmpl_selected_bd":  "#2D2D7A",
    "tmpl_hover_bg":     "#F5F6FD",
}

import platform as _platform
_OS = _platform.system()
_FONT_FAMILY = (
    "Segoe UI"        if _OS == "Windows" else
    "SF Pro Display"  if _OS == "Darwin"  else
    "Noto Sans"
)
_FONT_ARABIC = "Segoe UI" if _OS == "Windows" else "Arial"

def _F(size: int, weight: str = "normal") -> tuple:
    return (_FONT_FAMILY, size, weight)


if _TK_AVAILABLE:

    class _HoverButton(tk.Button):
        def __init__(self, master, bg_normal: str, bg_hover: str, **kwargs):
            super().__init__(master, bg=bg_normal, **kwargs)
            self._bg_n, self._bg_h = bg_normal, bg_hover
            self.bind("<Enter>", lambda e: self.config(bg=self._bg_h)
                      if str(self["state"]) != "disabled" else None)
            self.bind("<Leave>", lambda e: self.config(bg=self._bg_n)
                      if str(self["state"]) != "disabled" else None)


    class _FileCard(tk.Frame):
        """Styled drop-zone file selector card."""

        def __init__(self, master, label: str, icon: str, callback, **kwargs):
            super().__init__(master, bg=_C["card"],
                             highlightthickness=1,
                             highlightbackground=_C["border"], **kwargs)
            self._callback = callback
            self._path_var = tk.StringVar()

            top = tk.Frame(self, bg=_C["card"])
            top.pack(fill="x", padx=16, pady=(12, 0))
            tk.Label(top, text=icon, font=(_FONT_FAMILY, 18),
                     bg=_C["card"], fg=_C["navy"]).pack(side="left")
            tk.Label(top, text=label, font=_F(10, "bold"),
                     bg=_C["card"], fg=_C["text_h"]).pack(side="left", padx=8)

            self._zone = tk.Frame(self, bg=_C["bg"],
                                  highlightthickness=1,
                                  highlightbackground=_C["border"])
            self._zone.pack(fill="x", padx=16, pady=8)

            self._name_lbl = tk.Label(self._zone, text="No file selected",
                                      font=_F(9), bg=_C["bg"],
                                      fg=_C["text_muted"], anchor="w",
                                      wraplength=320)
            self._name_lbl.pack(side="left", fill="x", expand=True,
                                padx=12, pady=10)

            _HoverButton(self._zone,
                         bg_normal=_C["btn_secondary"],
                         bg_hover=_C["btn_secondary_hover"],
                         text="Browse", font=_F(9, "bold"),
                         fg=_C["btn_secondary_fg"],
                         relief="flat", cursor="hand2",
                         padx=14, pady=5,
                         command=self._callback).pack(side="right", padx=10, pady=8)

            for w in (self._zone, self._name_lbl):
                w.bind("<Enter>", lambda e: self._zone.config(
                    highlightbackground=_C["border_focus"]))
                w.bind("<Leave>", lambda e: self._zone.config(
                    highlightbackground=_C["border"] if not self._path_var.get()
                    else _C["green"]))

        def set_path(self, path: str):
            self._path_var.set(path)
            if path:
                self._name_lbl.config(text=os.path.basename(path), fg=_C["text_body"])
                self._zone.config(highlightbackground=_C["green"])
            else:
                self._name_lbl.config(text="No file selected", fg=_C["text_muted"])
                self._zone.config(highlightbackground=_C["border"])

        def get(self) -> str:
            return self._path_var.get()


    class _TemplateCard(tk.Frame):
        """Selectable card for one output template."""

        def __init__(self, master, key: str, info: Dict,
                     var: "tk.StringVar", **kwargs):
            super().__init__(master, bg=_C["card"],
                             highlightthickness=2,
                             highlightbackground=_C["border"],
                             cursor="hand2", **kwargs)
            self._key = key
            self._var = var
            self._info = info

            tk.Label(self, text=info["icon"], font=(_FONT_FAMILY, 22),
                     bg=_C["card"], fg=_C["navy"]).pack(side="left",
                                                         padx=(14, 8), pady=12)

            mid = tk.Frame(self, bg=_C["card"])
            mid.pack(side="left", fill="both", expand=True, pady=12)
            tk.Label(mid, text=info["label"], font=_F(10, "bold"),
                     bg=_C["card"], fg=_C["text_h"], anchor="w").pack(fill="x")
            tk.Label(mid, text=info["description"], font=_F(8),
                     bg=_C["card"], fg=_C["text_muted"],
                     anchor="w", wraplength=300).pack(fill="x")

            rb = tk.Radiobutton(self, variable=var, value=key,
                                bg=_C["card"], activebackground=_C["card"],
                                cursor="hand2",
                                command=self._on_select)
            rb.pack(side="right", padx=14)
            self._rb = rb

            for w in (self, mid) + tuple(mid.winfo_children()):
                w.bind("<Button-1>", lambda e: self._on_select())
                w.bind("<Enter>",    self._on_enter)
                w.bind("<Leave>",    self._on_leave)

            var.trace_add("write", lambda *a: self._refresh())
            self._refresh()

        def _on_select(self):
            self._var.set(self._key)

        def _on_enter(self, _=None):
            if self._var.get() != self._key:
                self.config(bg=_C["tmpl_hover_bg"],
                            highlightbackground=_C["border_focus"])
                for w in self.winfo_children():
                    try:
                        w.config(bg=_C["tmpl_hover_bg"])
                    except Exception:
                        pass

        def _on_leave(self, _=None):
            self._refresh()

        def _refresh(self):
            selected = (self._var.get() == self._key)
            bg = _C["tmpl_selected_bg"] if selected else _C["card"]
            bd = _C["tmpl_selected_bd"]  if selected else _C["border"]
            self.config(bg=bg, highlightbackground=bd)
            for w in self.winfo_children():
                try:
                    w.config(bg=bg)
                except Exception:
                    pass
            try:
                self._rb.config(bg=bg, activebackground=bg)
            except Exception:
                pass


    class _StatusBar(tk.Frame):
        """Animated progress stripe + colour-coded message."""

        def __init__(self, master, **kwargs):
            super().__init__(master, bg=_C["bg"], **kwargs)
            self._anim_id = None
            self._animating = False
            self._anim_pos = 0

            self._canvas = tk.Canvas(self, height=4, bg=_C["border"],
                                     highlightthickness=0)
            self._canvas.pack(fill="x")

            msg_row = tk.Frame(self, bg=_C["bg"])
            msg_row.pack(fill="x", pady=(4, 0))
            self._icon_lbl = tk.Label(msg_row, text="", font=_F(11),
                                      bg=_C["bg"], width=2)
            self._icon_lbl.pack(side="left", padx=(0, 4))
            self._msg_lbl = tk.Label(msg_row, text="", font=_F(9),
                                     bg=_C["bg"], fg=_C["text_muted"], anchor="w")
            self._msg_lbl.pack(side="left", fill="x", expand=True)

        def set(self, text: str, kind: str = "info"):
            colors = {
                "info":    (_C["info"],    "ℹ"),
                "success": (_C["success"], "✓"),
                "error":   (_C["error"],   "✗"),
                "warning": (_C["warning"], "⚠"),
            }
            fg, icon = colors.get(kind, (_C["text_muted"], ""))
            self._msg_lbl.config(text=text, fg=fg)
            self._icon_lbl.config(text=icon, fg=fg)

        def start_progress(self):
            if self._animating:
                return
            self._animating = True
            self._anim_pos = 0
            self._tick()

        def stop_progress(self):
            self._animating = False
            if self._anim_id:
                try:
                    self._canvas.after_cancel(self._anim_id)
                except Exception:
                    pass
            self._canvas.delete("all")

        def show_success_bar(self):
            self.stop_progress()
            w = self._canvas.winfo_width() or 600
            self._canvas.delete("all")
            self._canvas.create_rectangle(0, 0, w, 4, fill=_C["success"], outline="")

        def _tick(self):
            if not self._animating:
                return
            w = self._canvas.winfo_width() or 600
            self._canvas.delete("all")
            stripe_w = w // 3
            x = self._anim_pos % (w + stripe_w) - stripe_w
            self._canvas.create_rectangle(0, 0, w, 4, fill=_C["border"], outline="")
            self._canvas.create_rectangle(x, 0, x + stripe_w, 4,
                                           fill=_C["navy_mid"], outline="")
            self._anim_pos += 8
            self._anim_id = self._canvas.after(30, self._tick)


    class _LogPanel(tk.Frame):
        """Scrollable timestamped processing log."""

        def __init__(self, master, **kwargs):
            super().__init__(master, bg=_C["card"],
                             highlightthickness=1,
                             highlightbackground=_C["border"], **kwargs)
            hdr = tk.Frame(self, bg=_C["card"])
            hdr.pack(fill="x", padx=12, pady=(8, 4))
            tk.Label(hdr, text="Processing Log", font=_F(9, "bold"),
                     bg=_C["card"], fg=_C["text_muted"]).pack(side="left")
            tk.Button(hdr, text="Clear", font=_F(8), bg=_C["card"],
                      fg=_C["text_muted"], relief="flat", cursor="hand2",
                      command=self.clear).pack(side="right")

            self._text = tk.Text(self, height=5, font=(_FONT_ARABIC, 8),
                                 bg=_C["bg"], fg=_C["text_body"],
                                 relief="flat", wrap="word",
                                 state="disabled", padx=8, pady=4)
            sb = tk.Scrollbar(self, command=self._text.yview, width=10)
            self._text.config(yscrollcommand=sb.set)
            self._text.pack(side="left", fill="both", expand=True,
                            padx=(8, 0), pady=(0, 8))
            sb.pack(side="right", fill="y", pady=(0, 8), padx=(0, 8))
            self._text.tag_config("info",    foreground=_C["text_body"])
            self._text.tag_config("success", foreground=_C["success"])
            self._text.tag_config("error",   foreground=_C["error"])
            self._text.tag_config("warning", foreground=_C["warning"])

        def append(self, text: str, kind: str = "info"):
            import time
            ts = time.strftime("%H:%M:%S")
            self._text.config(state="normal")
            self._text.insert("end", f"[{ts}] {text}\n", kind)
            self._text.see("end")
            self._text.config(state="disabled")

        def clear(self):
            self._text.config(state="normal")
            self._text.delete("1.0", "end")
            self._text.config(state="disabled")


    class TransferApp:
        """ETEC MCQ Bank Transfer Tool — Enterprise GUI."""

        _W, _H = 660, 720

        def __init__(self, root: "tk.Tk"):
            self.root = root
            self._source_fmt    = tk.StringVar(value="auto")
            self._template_key  = tk.StringVar(value=_DEFAULT_TEMPLATE_KEY)

            self._setup_window()
            self._validate_templates()
            self._build_ui()

        def _setup_window(self):
            self.root.title("ETEC — MCQ Bank Transfer Tool")
            self.root.geometry(f"{self._W}x{self._H}")
            self.root.resizable(True, True)
            self.root.minsize(560, 640)
            self.root.configure(bg=_C["bg"])
            try:
                import base64
                from io import BytesIO
                from PIL import Image, ImageTk
                img  = Image.open(BytesIO(base64.b64decode(_ETEC_LOGO_B64)))
                icon = ImageTk.PhotoImage(img.resize((32, 32), Image.LANCZOS))
                self.root.iconphoto(True, icon)
                self._icon_ref = icon
            except Exception:
                pass

        def _validate_templates(self):
            results = template_manager.validate_all()
            ok  = [k for k, v in results.items() if v]
            bad = [k for k, v in results.items() if not v]
            if ok:
                log.info("Templates validated OK: %s", ok)
            if bad:
                log.warning("Template validation failures: %s", bad)
            self._template_validation = results

        def _build_ui(self):
            self._build_header()

            content = tk.Frame(self.root, bg=_C["bg"])
            content.pack(fill="both", expand=True, padx=20, pady=0)

            self._build_format_section(content)
            self._build_source_card(content)
            self._build_template_section(content)
            self._build_output_card(content)
            self._build_action_section(content)

            self._status_bar = _StatusBar(content)
            self._status_bar.pack(fill="x", pady=(8, 0))
            self._status_bar.set("Ready — select a source file to begin.", "info")

            self._log = _LogPanel(content)
            self._log.pack(fill="both", expand=True, pady=(8, 12))

            self.root.bind("<Return>", lambda e: self._start_transfer())
            self.root.bind("<Escape>", lambda e: self.root.focus())

        def _build_header(self):
            header = tk.Frame(self.root, bg=_C["navy_dark"], height=80)
            header.pack(fill="x")
            header.pack_propagate(False)

            try:
                import base64
                from io import BytesIO
                from PIL import Image, ImageTk
                raw     = base64.b64decode(_ETEC_LOGO_B64)
                pil_img = Image.open(BytesIO(raw)).convert("RGBA")
                bg_rgb  = tuple(int(_C["navy_dark"].lstrip("#")[i:i+2], 16)
                                for i in (0, 2, 4))
                bg_pil  = Image.new("RGBA", pil_img.size, bg_rgb + (255,))
                bg_pil.alpha_composite(pil_img)
                logo_pil = bg_pil.convert("RGB").resize((130, 50), Image.LANCZOS)
                self._logo_img = ImageTk.PhotoImage(logo_pil)
                tk.Label(header, image=self._logo_img,
                         bg=_C["navy_dark"], bd=0).pack(side="left",
                                                         padx=20, pady=15)
            except Exception:
                tk.Label(header, text="ETEC", font=_F(16, "bold"),
                         bg=_C["navy_dark"], fg="white").pack(side="left", padx=20)

            title_frame = tk.Frame(header, bg=_C["navy_dark"])
            title_frame.pack(side="right", padx=20)
            tk.Label(title_frame, text="MCQ Bank Transfer Tool",
                     font=_F(13, "bold"),
                     bg=_C["navy_dark"], fg="white").pack(anchor="e")
            tk.Label(title_frame, text="Question Bank Transformation Engine",
                     font=_F(8), bg=_C["navy_dark"],
                     fg="#8898C8").pack(anchor="e")

            tk.Frame(self.root, bg=_C["teal"], height=3).pack(fill="x")

        def _build_format_section(self, parent):
            row = tk.Frame(parent, bg=_C["bg"])
            row.pack(fill="x", pady=(12, 0))
            tk.Label(row, text="Source Format:",
                     font=_F(9, "bold"),
                     bg=_C["bg"], fg=_C["text_h"]).pack(side="left")
            for label, val in [("Auto-Detect","auto"),("PEA","pea"),
                                ("Nymr","nymr"),("Simple","simple")]:
                tk.Radiobutton(row, text=label,
                               variable=self._source_fmt, value=val,
                               font=_F(9), bg=_C["bg"], fg=_C["text_body"],
                               activebackground=_C["bg"],
                               selectcolor=_C["bg"],
                               cursor="hand2").pack(side="left", padx=(10, 0))
            self._detect_lbl = tk.Label(row, text="", font=_F(8, "italic"),
                                        bg=_C["bg"], fg=_C["teal"])
            self._detect_lbl.pack(side="left", padx=10)

        def _build_source_card(self, parent):
            self._src_card = _FileCard(parent,
                                       label="Source Questions File",
                                       icon="📂",
                                       callback=self._browse_source)
            self._src_card.pack(fill="x", pady=(12, 0))

        def _build_template_section(self, parent):
            hdr = tk.Frame(parent, bg=_C["bg"])
            hdr.pack(fill="x", pady=(14, 4))
            tk.Label(hdr, text="Output Template",
                     font=_F(10, "bold"),
                     bg=_C["bg"], fg=_C["text_h"]).pack(side="left")
            badge_text = f"{len(template_manager.available)} templates available"
            tk.Label(hdr, text=badge_text, font=_F(8),
                     bg=_C["bg"], fg=_C["text_muted"]).pack(side="left", padx=8)

            cards_frame = tk.Frame(parent, bg=_C["bg"])
            cards_frame.pack(fill="x")
            cards_frame.columnconfigure(0, weight=1)
            cards_frame.columnconfigure(1, weight=1)

            self._tmpl_cards = {}
            for col_idx, key in enumerate(template_manager.available):
                info = template_manager.get_info(key)
                card = _TemplateCard(cards_frame, key=key, info=info,
                                      var=self._template_key)
                card.grid(row=0, column=col_idx, padx=(0 if col_idx else 0, 6),
                          sticky="nsew")
                self._tmpl_cards[key] = card
                card.bind("<space>", lambda e, k=key: self._template_key.set(k))

            bad_tmpls = [k for k, ok in self._template_validation.items() if not ok]
            if bad_tmpls:
                warn = tk.Frame(parent, bg=_C["warning_bg"],
                                highlightthickness=1,
                                highlightbackground=_C["warning"])
                warn.pack(fill="x", pady=(6, 0))
                tk.Label(warn,
                         text=f"⚠  Templates with issues: {', '.join(bad_tmpls)}. "
                              f"Processing with these templates may fail.",
                         font=_F(8), bg=_C["warning_bg"], fg=_C["warning"],
                         wraplength=580, justify="left",
                         padx=10, pady=6).pack(fill="x")

        def _build_output_card(self, parent):
            self._out_card = _FileCard(parent,
                                       label="Output File  (Save As)",
                                       icon="💾",
                                       callback=self._browse_output)
            self._out_card.pack(fill="x", pady=(10, 0))

        def _build_action_section(self, parent):
            row = tk.Frame(parent, bg=_C["bg"])
            row.pack(fill="x", pady=(16, 0))

            self._run_btn = _HoverButton(
                row,
                bg_normal=_C["btn_primary"],
                bg_hover=_C["btn_primary_hover"],
                text="  ▶  Process Questions",
                font=_F(11, "bold"),
                fg=_C["btn_primary_fg"],
                relief="flat", cursor="hand2",
                padx=28, pady=12,
                command=self._start_transfer,
            )
            self._run_btn.pack(side="left", expand=True, fill="x")

            _HoverButton(
                row,
                bg_normal=_C["btn_secondary"],
                bg_hover=_C["btn_secondary_hover"],
                text="↺  Reset",
                font=_F(9),
                fg=_C["btn_secondary_fg"],
                relief="flat", cursor="hand2",
                padx=16, pady=12,
                command=self._reset,
            ).pack(side="right", padx=(8, 0))

        def _browse_source(self):
            path = filedialog.askopenfilename(
                title="Select Source Questions File",
                filetypes=[("Word Documents", "*.docx"), ("All Files", "*.*")],
            )
            if not path:
                return
            self._src_card.set_path(path)
            self._detect_lbl.config(text="")
            if self._source_fmt.get() == "auto":
                try:
                    doc = docx.Document(path)
                    fmt = detect_source_format(doc)
                    labels = {"pea": "Detected: PEA",
                              "nymr": "Detected: Nymr",
                              "simple": "Detected: Simple"}
                    self._detect_lbl.config(text=labels.get(fmt, f"Detected: {fmt}"))
                    self._log.append(f"Source loaded — format: {fmt.upper()}", "info")
                except Exception as e:
                    self._detect_lbl.config(text="Detection failed")
                    self._log.append(f"Format detection error: {e}", "warning")

        def _browse_output(self):
            path = filedialog.asksaveasfilename(
                title="Save Output As",
                defaultextension=".docx",
                filetypes=[("Word Documents", "*.docx")],
                initialfile="Final_Question_Bank.docx",
            )
            if path:
                self._out_card.set_path(path)
                self._log.append(f"Output: {path}", "info")

        def _start_transfer(self):
            src  = self._src_card.get()
            out  = self._out_card.get()
            key  = self._template_key.get()

            if not src:
                messagebox.showwarning("Missing Input",
                                       "Please select a source questions file.")
                return
            if not out:
                messagebox.showwarning("Missing Output",
                                       "Please choose where to save the output file.")
                return
            if not key or key not in template_manager.available:
                messagebox.showwarning("No Template Selected",
                                       "Please select an output template.")
                return

            tmpl_info = template_manager.get_info(key)
            self._run_btn.config(text="⏳  Processing…",
                                 state=tk.DISABLED, bg=_C["btn_disabled"])
            self._status_bar.set("Processing — please wait…", "info")
            self._status_bar.start_progress()
            self._log.append(f"Starting transfer…", "info")
            self._log.append(f"Source:   {src}", "info")
            self._log.append(f"Template: {tmpl_info['label']} (key={key})", "info")
            self._log.append(f"Output:   {out}", "info")

            threading.Thread(target=self._run_backend,
                             args=(src, out, key), daemon=True).start()

        def _run_backend(self, source: str, output: str, template_key: str):
            import traceback

            def _ui(fn):
                self.root.after(0, fn)

            try:
                process_and_transfer(source, output, template_key=template_key)

                def _success():
                    self._status_bar.stop_progress()
                    self._status_bar.show_success_bar()
                    self._status_bar.set("Complete — output saved successfully.", "success")
                    self._log.append("Transfer completed successfully.", "success")
                    self._run_btn.config(text="✓  Done",
                                         state=tk.NORMAL, bg=_C["success"])
                    messagebox.showinfo("Transfer Complete",
                                        f"✓  Question bank generated successfully."
                                        f"\n\nSaved to:\n{output}")
                _ui(_success)

            except Exception as exc:
                tb = traceback.format_exc()
                def _error():
                    self._status_bar.stop_progress()
                    self._status_bar.set("Transfer failed — see log for details.", "error")
                    self._log.append(f"ERROR: {exc}", "error")
                    for line in tb.strip().splitlines()[-6:]:
                        self._log.append(f"  {line}", "error")
                    self._run_btn.config(text="  ▶  Process Questions",
                                         state=tk.NORMAL, bg=_C["btn_primary"])
                    messagebox.showerror("Transfer Failed",
                                         f"An error occurred during processing:\n\n{exc}")
                _ui(_error)

        def _reset(self):
            self._src_card.set_path("")
            self._out_card.set_path("")
            self._detect_lbl.config(text="")
            self._status_bar.set("Ready — select a source file to begin.", "info")
            self._status_bar.stop_progress()
            self._run_btn.config(text="  ▶  Process Questions",
                                 state=tk.NORMAL, bg=_C["btn_primary"])
            self._log.append("Reset.", "info")


# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
# ENTRY POINT
# ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
if __name__ == "__main__":
    if not _TK_AVAILABLE:
        raise SystemExit(
            "tkinter is not available in this environment; the GUI cannot start. "
            "The core transfer functions (process_and_transfer, parse_source, "
            "inject_images, ...) remain importable and usable as a library."
        )
    root = tk.Tk()
    TransferApp(root)
    root.mainloop()
